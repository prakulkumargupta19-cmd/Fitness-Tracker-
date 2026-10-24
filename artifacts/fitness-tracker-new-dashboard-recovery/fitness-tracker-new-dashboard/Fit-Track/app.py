from flask import Flask, render_template, request, redirect, url_for, flash, jsonify, session, abort, send_from_directory, Response
import sqlite3, os, csv, io, json, uuid
from datetime import datetime, date, timedelta
from functools import wraps
from werkzeug.security import generate_password_hash, check_password_hash
from werkzeug.utils import secure_filename

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DB_PATH = os.path.join(BASE_DIR, "data", "sports.db")
app = Flask(__name__)
app.secret_key = os.environ.get("SPORTS_PASSPORT_SECRET", "change-this-before-production")

SPORTS = [
    ("Football", "Coached", "Football Ground"),
    ("Futsal", "Activity", "Futsal Arena"),
    ("Basketball", "Coached", "Basketball Court"),
    ("Gym", "Coached", "Gym & Fitness Centre"),
    ("Box Cricket", "Activity", "Box Cricket Arena"),
    ("Cricket Nets", "Activity", "Cricket Practice Nets"),
    ("Volleyball", "Activity", "Volleyball Court"),
    ("Badminton", "Activity", "Indoor Sports Arena"),
    ("Pickleball", "Activity", "Pickleball Courts"),
    ("Pool", "Activity", "Indoor Games Area"),
    ("Carrom", "Activity", "Indoor Games Area"),
    ("Foosball", "Activity", "Indoor Games Area"),
    ("Table Tennis", "Activity", "Indoor Sports Arena"),
    ("Kick Boxing", "Coached", "Combat/Fitness Area"),
]

COACHES = [
    ("Rajendra Rajat", "Basketball", "rajendra.rajat"),
    ("Nagen Tamang", "Football", "nagen.tamang"),
    ("Ritik Jaiswal", "Gym", "ritik.jaiswal"),
    ("Siva Gurung", "Kick Boxing", "siva.gurung"),
]

ADMINS = [
    ("Mr. Sushovan Gupta", "HOD, BBA Sports Management", "sushovan.gupta"),
    ("Mr. Neehar Reddy", "HOS, BBA Sports Management", "neehar.reddy"),
]

DEPARTMENTS = [
    "BBA", "BBA DM", "BBA HM", "BBA SM", "BBA ATA", "BBA E", "BBA GB",
    "BMAGD + BMS + BID + BFDM", "BCA", "BSC DA + BSC AI + BCA AI&ML + BCA DS&CS", "HHA + BAHS"
]

DEPARTMENT_ALIASES = {
    "bba sports management": "bba sm",
    "bba digital marketing": "bba dm",
    "bba hospital management": "bba hm",
    "bba entrepreneurship": "bba e",
    "bba global business": "bba gb",
    "bba aviation tourism & aviation": "bba ata",
}

def department_match(student_department, team_department):
    left=(student_department or "").strip().lower(); right=(team_department or "").strip().lower()
    return left == right or DEPARTMENT_ALIASES.get(left) == right

LIVE_SCORE_FIELDS = {
    "Football": [("period", "Half / Period"), ("possession", "Possession"), ("corners", "Corners"), ("fouls", "Fouls"), ("yellow_cards", "Yellow Cards"), ("red_cards", "Red Cards")],
    "Futsal": [("period", "Half / Period"), ("fouls", "Team Fouls"), ("timeouts", "Timeouts"), ("yellow_cards", "Yellow Cards"), ("red_cards", "Red Cards")],
    "Basketball": [("quarter", "Quarter"), ("clock", "Game Clock"), ("team_fouls", "Team Fouls"), ("timeouts", "Timeouts"), ("possession", "Possession Arrow")],
    "Cricket": [("overs", "Overs"), ("wickets", "Wickets"), ("run_rate", "Run Rate"), ("target", "Target"), ("striker", "Striker")],
    "Box Cricket": [("overs", "Overs"), ("wickets", "Wickets"), ("run_rate", "Run Rate"), ("target", "Target")],
    "Volleyball": [("set", "Current Set"), ("sets", "Sets Won"), ("timeouts", "Timeouts")],
    "Table Tennis": [("game", "Current Game"), ("games", "Games Won"), ("serve", "Serving")],
    "Badminton": [("game", "Current Game"), ("games", "Games Won"), ("serve", "Serving")],
    "Pickleball": [("game", "Current Game"), ("games", "Games Won"), ("serve", "Serving")],
}

UPLOAD_ROOT = os.path.join(BASE_DIR, "static", "uploads")
ALLOWED_IMAGE_EXTENSIONS = {"png", "jpg", "jpeg", "webp"}
ALLOWED_VIDEO_EXTENSIONS = {"mp4", "webm", "mov"}
ALLOWED_DRAW_EXTENSIONS = {"pdf"}
ALLOWED_CERTIFICATE_EXTENSIONS = {"pdf", "png", "jpg", "jpeg"}
ALLOWED_PARTICIPANT_SHEETS = {"xlsx", "xls"}

SCORE_FIELDS = ["attendance_score", "skill_development", "passing", "dribbling", "shooting", "vision", "fitness", "overall_performance"]

SPORT_METRICS = {
    "Football": [
        ("skill_development", "Skill Development"), ("ball_control", "Ball Control"),
        ("passing", "Passing"), ("dribbling", "Dribbling"), ("shooting", "Shooting"),
        ("vision", "Vision"), ("tactical_awareness", "Tactical Awareness"),
        ("fitness", "Fitness"), ("work_rate", "Work Rate"),
    ],
    "Basketball": [
        ("skill_development", "Skill Development"), ("ball_handling", "Ball Handling"),
        ("passing", "Passing"), ("dribbling", "Dribbling"), ("shooting", "Shooting"),
        ("court_vision", "Court Vision"), ("defence", "Defence"),
        ("fitness", "Fitness"), ("team_play", "Team Play"),
    ],
    "Gym": [
        ("strength", "Strength"), ("endurance", "Endurance"), ("mobility", "Mobility"),
        ("technique_form", "Technique / Form"), ("consistency", "Consistency"),
        ("discipline", "Discipline"), ("recovery", "Recovery"), ("fitness", "Fitness"),
    ],
    "Kick Boxing": [
        ("stance_guard", "Stance & Guard"), ("footwork", "Footwork"),
        ("punching", "Punching Technique"), ("kicking", "Kicking Technique"),
        ("defence", "Defence"), ("speed_timing", "Speed & Timing"),
        ("fitness", "Fitness"), ("discipline", "Discipline"),
    ],
}

ATTENDANCE_SCORES = {"Present": 10, "Late": 7, "Absent": 1, "Excused": None}


def get_db():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys = ON")
    return conn


def table_columns(conn, table):
    return {r["name"] for r in conn.execute(f"PRAGMA table_info({table})").fetchall()}


def ensure_column(conn, table, column_definition):
    name = column_definition.split()[0]
    if name not in table_columns(conn, table):
        conn.execute(f"ALTER TABLE {table} ADD COLUMN {column_definition}")




def extend_v5_schema(conn):
    """V5 operational schema: official participants, fixture rosters, results, draws, announcements, certificates and facility galleries."""
    ensure_column(conn, "students", "passport_id TEXT")
    ensure_column(conn, "students", "photo_path TEXT")
    ensure_column(conn, "coaches", "photo_path TEXT")
    ensure_column(conn, "matches", "category TEXT DEFAULT 'Boys'")
    ensure_column(conn, "sports", "cover_path TEXT")
    conn.executescript("""
    CREATE UNIQUE INDEX IF NOT EXISTS idx_students_passport_id ON students(passport_id) WHERE passport_id IS NOT NULL;
    CREATE TABLE IF NOT EXISTS event_participants (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        tournament_id INTEGER NOT NULL,
        student_id INTEGER,
        name TEXT NOT NULL,
        department TEXT NOT NULL,
        category TEXT NOT NULL,
        year TEXT NOT NULL,
        imported_at TEXT DEFAULT CURRENT_TIMESTAMP,
        UNIQUE(tournament_id,name,department,category,year),
        FOREIGN KEY(tournament_id) REFERENCES tournaments(id) ON DELETE CASCADE,
        FOREIGN KEY(student_id) REFERENCES students(id) ON DELETE SET NULL
    );
    CREATE TABLE IF NOT EXISTS match_players (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        match_id INTEGER NOT NULL,
        side TEXT NOT NULL CHECK(side IN ('A','B')),
        student_name TEXT NOT NULL,
        department TEXT,
        student_id INTEGER,
        FOREIGN KEY(match_id) REFERENCES matches(id) ON DELETE CASCADE,
        FOREIGN KEY(student_id) REFERENCES students(id) ON DELETE SET NULL
    );
    CREATE TABLE IF NOT EXISTS tournament_results (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        tournament_id INTEGER NOT NULL,
        category TEXT NOT NULL,
        winner TEXT NOT NULL,
        runner_up TEXT NOT NULL,
        notes TEXT,
        UNIQUE(tournament_id,category),
        FOREIGN KEY(tournament_id) REFERENCES tournaments(id) ON DELETE CASCADE
    );
    CREATE TABLE IF NOT EXISTS draw_documents (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        tournament_id INTEGER NOT NULL,
        category TEXT NOT NULL,
        format_type TEXT NOT NULL,
        file_path TEXT NOT NULL,
        uploaded_by INTEGER,
        created_at TEXT DEFAULT CURRENT_TIMESTAMP,
        FOREIGN KEY(tournament_id) REFERENCES tournaments(id) ON DELETE CASCADE,
        FOREIGN KEY(uploaded_by) REFERENCES users(id) ON DELETE SET NULL
    );
    CREATE TABLE IF NOT EXISTS announcements (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        title TEXT NOT NULL,
        message TEXT NOT NULL,
        created_by INTEGER,
        active INTEGER DEFAULT 1,
        created_at TEXT DEFAULT CURRENT_TIMESTAMP,
        FOREIGN KEY(created_by) REFERENCES users(id) ON DELETE SET NULL
    );
    CREATE TABLE IF NOT EXISTS certificate_templates (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        tournament_id INTEGER NOT NULL,
        certificate_type TEXT NOT NULL,
        file_path TEXT NOT NULL,
        uploaded_by INTEGER,
        created_at TEXT DEFAULT CURRENT_TIMESTAMP,
        UNIQUE(tournament_id,certificate_type),
        FOREIGN KEY(tournament_id) REFERENCES tournaments(id) ON DELETE CASCADE,
        FOREIGN KEY(uploaded_by) REFERENCES users(id) ON DELETE SET NULL
    );
    CREATE TABLE IF NOT EXISTS generated_certificates (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        tournament_id INTEGER NOT NULL,
        participant_id INTEGER,
        student_id INTEGER,
        certificate_type TEXT NOT NULL,
        recipient_name TEXT NOT NULL,
        achievement TEXT,
        certificate_no TEXT UNIQUE NOT NULL,
        file_path TEXT NOT NULL,
        created_at TEXT DEFAULT CURRENT_TIMESTAMP,
        FOREIGN KEY(tournament_id) REFERENCES tournaments(id) ON DELETE CASCADE,
        FOREIGN KEY(participant_id) REFERENCES event_participants(id) ON DELETE SET NULL,
        FOREIGN KEY(student_id) REFERENCES students(id) ON DELETE SET NULL
    );
    CREATE TABLE IF NOT EXISTS facility_media (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        sport_id INTEGER NOT NULL,
        file_path TEXT NOT NULL,
        caption TEXT,
        uploaded_by INTEGER,
        created_at TEXT DEFAULT CURRENT_TIMESTAMP,
        FOREIGN KEY(sport_id) REFERENCES sports(id) ON DELETE CASCADE,
        FOREIGN KEY(uploaded_by) REFERENCES users(id) ON DELETE SET NULL
    );
    """)
    ensure_column(conn, "announcements", "image_path TEXT")
    # Give every existing student a permanent Sports Passport ID.
    rows = conn.execute("SELECT id,passport_id FROM students ORDER BY id").fetchall()
    for row in rows:
        if not row["passport_id"]:
            conn.execute("UPDATE students SET passport_id=? WHERE id=?", (make_passport_id(row["id"]), row["id"]))


def extend_v7_schema(conn):
    """V7 coaching schedule schema. Non-destructive upgrade for existing databases."""
    conn.executescript("""
    CREATE TABLE IF NOT EXISTS coaching_sessions (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        programme_id INTEGER NOT NULL,
        session_date TEXT NOT NULL,
        start_time TEXT NOT NULL,
        end_time TEXT NOT NULL,
        venue TEXT,
        notes TEXT,
        status TEXT DEFAULT 'Scheduled',
        created_by INTEGER,
        updated_by INTEGER,
        created_at TEXT DEFAULT CURRENT_TIMESTAMP,
        updated_at TEXT DEFAULT CURRENT_TIMESTAMP,
        FOREIGN KEY(programme_id) REFERENCES programmes(id) ON DELETE CASCADE,
        FOREIGN KEY(created_by) REFERENCES users(id) ON DELETE SET NULL,
        FOREIGN KEY(updated_by) REFERENCES users(id) ON DELETE SET NULL
    );
    CREATE INDEX IF NOT EXISTS idx_coaching_sessions_date ON coaching_sessions(session_date,start_time);
    CREATE INDEX IF NOT EXISTS idx_coaching_sessions_programme ON coaching_sessions(programme_id,session_date);
    """)


def can_edit_programme(conn, programme_id, user):
    if not user or user["role"] not in ("admin", "coach"):
        return False
    row = conn.execute("SELECT coach_id FROM programmes WHERE id=? AND active=1", (programme_id,)).fetchone()
    if not row:
        return False
    return user["role"] == "admin" or row["coach_id"] == user["coach_id"]


def make_passport_id(student_db_id):
    return f"IKC-SP-{datetime.now().year}-{int(student_db_id):05d}"


def next_certificate_no(tournament_id, certificate_type, serial):
    prefix = "ACH" if certificate_type.lower().startswith("achievement") else "PAR"
    return f"IKC-{datetime.now().year}-{int(tournament_id):03d}-{prefix}-{int(serial):04d}"


def participant_sheet_rows(file_storage):
    """Read XLSX/XLS official participant import. Required visible columns: Name, Department, Category, Year."""
    if not file_storage or not file_storage.filename or "." not in file_storage.filename:
        raise ValueError("Choose an XLSX or XLS participant file.")
    ext = file_storage.filename.rsplit(".", 1)[1].lower()
    if ext not in ALLOWED_PARTICIPANT_SHEETS:
        raise ValueError("Participant import supports .xlsx and .xls files only.")
    data = file_storage.read()
    raw_rows = []
    if ext == "xlsx":
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise ValueError("XLSX support is not installed. Run pip install -r requirements.txt.") from exc
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        ws = wb.active
        raw_rows = [list(r) for r in ws.iter_rows(values_only=True)]
    else:
        try:
            import xlrd
        except ImportError as exc:
            raise ValueError("XLS support requires xlrd. Run pip install -r requirements.txt.") from exc
        book = xlrd.open_workbook(file_contents=data)
        sh = book.sheet_by_index(0)
        raw_rows = [sh.row_values(i) for i in range(sh.nrows)]
    if not raw_rows:
        raise ValueError("The participant sheet is empty.")
    headers = [str(v or "").strip().lower().replace("_", " ") for v in raw_rows[0]]
    aliases = {
        "name": ["name", "student name", "participant name"],
        "department": ["department", "dept"],
        "category": ["category", "gender category", "boys/girls"],
        "year": ["year", "academic year", "batch/year"],
    }
    idx = {}
    for key, options in aliases.items():
        for option in options:
            if option in headers:
                idx[key] = headers.index(option); break
    missing = [k.title() for k in aliases if k not in idx]
    if missing:
        raise ValueError("Missing required participant columns: " + ", ".join(missing))
    out = []
    for row in raw_rows[1:]:
        vals = {k: str(row[i] if i < len(row) and row[i] is not None else "").strip() for k, i in idx.items()}
        if not vals["name"]:
            continue
        cat = vals["category"].title()
        if cat not in ("Boys", "Girls"):
            raise ValueError(f"Invalid category for {vals['name']}: use Boys or Girls.")
        vals["category"] = cat
        out.append(vals)
    return out


def create_certificate_pdf(template_full_path, output_full_path, recipient_name, event_name, category, year, certificate_type, certificate_no, achievement=None):
    """Create a downloadable certificate by overlaying data on an uploaded PDF/image template."""
    from reportlab.pdfgen import canvas
    from reportlab.lib.pagesizes import landscape, A4
    from reportlab.lib.utils import ImageReader
    from pypdf import PdfReader, PdfWriter
    ext = template_full_path.rsplit(".", 1)[-1].lower()
    os.makedirs(os.path.dirname(output_full_path), exist_ok=True)
    if ext == "pdf":
        template = PdfReader(template_full_path)
        if not template.pages:
            raise ValueError("Certificate template PDF has no pages.")
        page = template.pages[0]
        width = float(page.mediabox.width); height = float(page.mediabox.height)
        overlay_stream = io.BytesIO()
        c = canvas.Canvas(overlay_stream, pagesize=(width, height))
        c.setFillColorRGB(0.02, 0.12, 0.35)
        c.setFont("Helvetica-Bold", min(28, width/25))
        c.drawCentredString(width/2, height*0.49, recipient_name)
        c.setFont("Helvetica", min(13, width/55))
        detail = f"{certificate_type} - {event_name} - {category} - {year}"
        c.drawCentredString(width/2, height*0.41, detail)
        if achievement:
            c.setFont("Helvetica-Bold", min(15, width/48)); c.drawCentredString(width/2, height*0.35, achievement)
        c.setFont("Helvetica", min(9, width/75)); c.drawCentredString(width/2, height*0.10, f"Certificate No: {certificate_no}")
        c.save(); overlay_stream.seek(0)
        overlay = PdfReader(overlay_stream).pages[0]
        page.merge_page(overlay)
        writer = PdfWriter(); writer.add_page(page)
        with open(output_full_path, "wb") as f: writer.write(f)
    else:
        from PIL import Image
        img = Image.open(template_full_path)
        width_px, height_px = img.size
        width = landscape(A4)[0]; height = width * (height_px / width_px)
        c = canvas.Canvas(output_full_path, pagesize=(width, height))
        c.drawImage(ImageReader(img), 0, 0, width=width, height=height, preserveAspectRatio=True, mask='auto')
        c.setFillColorRGB(0.02, 0.12, 0.35)
        c.setFont("Helvetica-Bold", 26); c.drawCentredString(width/2, height*0.49, recipient_name)
        c.setFont("Helvetica", 12); c.drawCentredString(width/2, height*0.41, f"{certificate_type} - {event_name} - {category} - {year}")
        if achievement:
            c.setFont("Helvetica-Bold", 14); c.drawCentredString(width/2, height*0.35, achievement)
        c.setFont("Helvetica", 8); c.drawCentredString(width/2, height*0.09, f"Certificate No: {certificate_no}")
        c.save()


def init_db():
    os.makedirs(os.path.dirname(DB_PATH), exist_ok=True)
    conn = get_db()
    conn.executescript("""
    CREATE TABLE IF NOT EXISTS sports (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        name TEXT UNIQUE NOT NULL,
        activity_type TEXT NOT NULL DEFAULT 'Activity',
        venue TEXT,
        is_active INTEGER DEFAULT 1
    );
    CREATE TABLE IF NOT EXISTS students (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        name TEXT NOT NULL,
        student_id TEXT UNIQUE NOT NULL,
        email TEXT,
        phone TEXT,
        department TEXT,
        semester TEXT,
        sports TEXT,
        bio TEXT,
        emergency_contact TEXT,
        blood_group TEXT,
        medical_history TEXT,
        address TEXT,
        parent_name TEXT,
        parent_phone TEXT,
        parent_address TEXT,
        created_at TEXT DEFAULT CURRENT_TIMESTAMP
    );
    CREATE TABLE IF NOT EXISTS coaches (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        name TEXT NOT NULL,
        sport TEXT NOT NULL,
        email TEXT,
        username TEXT UNIQUE
    );
    CREATE TABLE IF NOT EXISTS users (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        username TEXT UNIQUE NOT NULL,
        password_hash TEXT NOT NULL,
        display_name TEXT NOT NULL,
        designation TEXT,
        role TEXT NOT NULL,
        coach_id INTEGER,
        student_id INTEGER,
        active INTEGER DEFAULT 1,
        FOREIGN KEY(coach_id) REFERENCES coaches(id) ON DELETE SET NULL,
        FOREIGN KEY(student_id) REFERENCES students(id) ON DELETE SET NULL
    );
    CREATE TABLE IF NOT EXISTS programmes (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        name TEXT NOT NULL,
        sport TEXT NOT NULL,
        coach_id INTEGER,
        schedule TEXT,
        venue TEXT,
        active INTEGER DEFAULT 1,
        FOREIGN KEY(coach_id) REFERENCES coaches(id) ON DELETE SET NULL
    );
    CREATE TABLE IF NOT EXISTS enrollments (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        student_id INTEGER NOT NULL,
        programme_id INTEGER NOT NULL,
        attendance INTEGER DEFAULT 0,
        total_sessions INTEGER DEFAULT 0,
        joined_on TEXT DEFAULT CURRENT_DATE,
        status TEXT DEFAULT 'Active',
        UNIQUE(student_id, programme_id),
        FOREIGN KEY(student_id) REFERENCES students(id) ON DELETE CASCADE,
        FOREIGN KEY(programme_id) REFERENCES programmes(id) ON DELETE CASCADE
    );
    CREATE TABLE IF NOT EXISTS attendance (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        enrollment_id INTEGER NOT NULL,
        session_date TEXT NOT NULL,
        present INTEGER DEFAULT 1,
        UNIQUE(enrollment_id, session_date),
        FOREIGN KEY(enrollment_id) REFERENCES enrollments(id) ON DELETE CASCADE
    );
    CREATE TABLE IF NOT EXISTS daily_assessments (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        enrollment_id INTEGER NOT NULL,
        coach_id INTEGER,
        session_date TEXT NOT NULL,
        attendance_status TEXT DEFAULT 'Present',
        attendance_score INTEGER,
        skill_development INTEGER,
        passing INTEGER,
        dribbling INTEGER,
        shooting INTEGER,
        vision INTEGER,
        fitness INTEGER,
        overall_performance INTEGER,
        remarks TEXT,
        created_at TEXT DEFAULT CURRENT_TIMESTAMP,
        updated_at TEXT DEFAULT CURRENT_TIMESTAMP,
        UNIQUE(enrollment_id, session_date),
        FOREIGN KEY(enrollment_id) REFERENCES enrollments(id) ON DELETE CASCADE,
        FOREIGN KEY(coach_id) REFERENCES coaches(id) ON DELETE SET NULL
    );
    CREATE TABLE IF NOT EXISTS assessment_metric_scores (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        assessment_id INTEGER NOT NULL,
        metric_key TEXT NOT NULL,
        metric_label TEXT NOT NULL,
        score INTEGER NOT NULL CHECK(score BETWEEN 1 AND 10),
        UNIQUE(assessment_id, metric_key),
        FOREIGN KEY(assessment_id) REFERENCES daily_assessments(id) ON DELETE CASCADE
    );
    CREATE TABLE IF NOT EXISTS tournaments (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        name TEXT NOT NULL,
        sport TEXT NOT NULL,
        start_date TEXT NOT NULL,
        end_date TEXT,
        venue TEXT,
        status TEXT DEFAULT 'Upcoming',
        description TEXT,
        poster_path TEXT,
        registration_open INTEGER DEFAULT 1
    );
    CREATE TABLE IF NOT EXISTS registrations (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        tournament_id INTEGER NOT NULL,
        student_id INTEGER NOT NULL,
        team_name TEXT,
        name TEXT,
        department TEXT,
        semester TEXT,
        phone TEXT,
        emergency_contact TEXT,
        blood_group TEXT,
        medical_history TEXT,
        registration_status TEXT DEFAULT 'Registered',
        registered_at TEXT DEFAULT CURRENT_TIMESTAMP,
        UNIQUE(tournament_id, student_id),
        FOREIGN KEY(tournament_id) REFERENCES tournaments(id) ON DELETE CASCADE,
        FOREIGN KEY(student_id) REFERENCES students(id) ON DELETE CASCADE
    );
    CREATE TABLE IF NOT EXISTS matches (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        tournament_id INTEGER NOT NULL,
        round_name TEXT,
        match_date TEXT,
        venue TEXT,
        team_a TEXT NOT NULL,
        team_b TEXT NOT NULL,
        score_a TEXT DEFAULT '0',
        score_b TEXT DEFAULT '0',
        status TEXT DEFAULT 'Scheduled',
        photo_path TEXT,
        live_note TEXT,
        FOREIGN KEY(tournament_id) REFERENCES tournaments(id) ON DELETE CASCADE
    );
    CREATE TABLE IF NOT EXISTS achievements (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        student_id INTEGER NOT NULL,
        title TEXT NOT NULL,
        sport TEXT,
        level TEXT,
        year TEXT,
        description TEXT,
        FOREIGN KEY(student_id) REFERENCES students(id) ON DELETE CASCADE
    );
    CREATE TABLE IF NOT EXISTS certificates (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        student_id INTEGER NOT NULL,
        title TEXT NOT NULL,
        issuer TEXT,
        issue_date TEXT,
        certificate_no TEXT,
        FOREIGN KEY(student_id) REFERENCES students(id) ON DELETE CASCADE
    );
    CREATE TABLE IF NOT EXISTS messages (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        sender_user_id INTEGER NOT NULL,
        receiver_user_id INTEGER NOT NULL,
        message TEXT NOT NULL,
        created_at TEXT DEFAULT CURRENT_TIMESTAMP,
        read_at TEXT,
        FOREIGN KEY(sender_user_id) REFERENCES users(id) ON DELETE CASCADE,
        FOREIGN KEY(receiver_user_id) REFERENCES users(id) ON DELETE CASCADE
    );
    CREATE TABLE IF NOT EXISTS event_media (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        tournament_id INTEGER,
        media_type TEXT NOT NULL DEFAULT 'photo',
        file_path TEXT NOT NULL,
        caption TEXT,
        uploaded_by INTEGER,
        created_at TEXT DEFAULT CURRENT_TIMESTAMP,
        FOREIGN KEY(tournament_id) REFERENCES tournaments(id) ON DELETE SET NULL,
        FOREIGN KEY(uploaded_by) REFERENCES users(id) ON DELETE SET NULL
    );
    CREATE TABLE IF NOT EXISTS match_live_stats (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        match_id INTEGER NOT NULL,
        stat_key TEXT NOT NULL,
        stat_label TEXT NOT NULL,
        stat_value TEXT,
        UNIQUE(match_id, stat_key),
        FOREIGN KEY(match_id) REFERENCES matches(id) ON DELETE CASCADE
    );
    """)
    # Safe upgrades for older project databases.
    ensure_column(conn, "coaches", "username TEXT")
    ensure_column(conn, "programmes", "active INTEGER DEFAULT 1")
    ensure_column(conn, "enrollments", "status TEXT DEFAULT 'Active'")
    for col in ["emergency_contact TEXT", "blood_group TEXT", "medical_history TEXT", "address TEXT", "parent_name TEXT", "parent_phone TEXT", "parent_address TEXT"]:
        ensure_column(conn, "students", col)
    ensure_column(conn, "users", "student_id INTEGER")
    for col in ["poster_path TEXT", "registration_open INTEGER DEFAULT 1"]:
        ensure_column(conn, "tournaments", col)
    for col in ["name TEXT", "department TEXT", "semester TEXT", "phone TEXT", "emergency_contact TEXT", "blood_group TEXT", "medical_history TEXT", "registration_status TEXT DEFAULT 'Registered'"]:
        ensure_column(conn, "registrations", col)
    for col in ["photo_path TEXT", "live_note TEXT"]:
        ensure_column(conn, "matches", col)
    extend_v5_schema(conn)
    extend_v7_schema(conn)
    seed(conn)
    extend_v5_schema(conn)
    extend_v7_schema(conn)
    migrate_legacy_metric_scores(conn)
    conn.commit()
    conn.close()


def seed(conn):
    for name, activity_type, venue in SPORTS:
        conn.execute("INSERT OR IGNORE INTO sports(name,activity_type,venue,is_active) VALUES (?,?,?,1)", (name, activity_type, venue))

    # Replace the original demo coach list with the Inspiria coaching team.
    for name, sport, username in COACHES:
        row = conn.execute("SELECT id FROM coaches WHERE name=?", (name,)).fetchone()
        if row:
            conn.execute("UPDATE coaches SET sport=?, username=? WHERE id=?", (sport, username, row["id"]))
        else:
            conn.execute("INSERT INTO coaches(name,sport,username) VALUES (?,?,?)", (name, sport, username))

    # Seed accounts. Default password is only for this local prototype.
    default_hash = generate_password_hash("Inspiria@2026")
    for name, designation, username in ADMINS:
        conn.execute("""INSERT OR IGNORE INTO users(username,password_hash,display_name,designation,role)
                        VALUES (?,?,?,?, 'admin')""", (username, default_hash, name, designation))
    for name, sport, username in COACHES:
        coach = conn.execute("SELECT id FROM coaches WHERE name=?", (name,)).fetchone()
        conn.execute("""INSERT OR IGNORE INTO users(username,password_hash,display_name,designation,role,coach_id)
                        VALUES (?,?,?,?, 'coach',?)""", (username, default_hash, name, f"{sport} Coach", coach["id"]))

    # Ensure the four coached programmes exist and point to the correct coach.
    programme_seed = [
        ("Basketball Coaching", "Basketball", "Rajendra Rajat", "Tue & Thu · 4:30 PM", "Basketball Court"),
        ("Inspiria Football Coaching", "Football", "Nagen Tamang", "Mon, Wed & Fri · 5:00 PM", "Football Ground"),
        ("Gym & Fitness Coaching", "Gym", "Ritik Jaiswal", "Weekdays · 6:00 AM & 4:00 PM", "Gym & Fitness Centre"),
        ("Kick Boxing Coaching", "Kick Boxing", "Siva Gurung", "Tue, Thu & Sat · 4:00 PM", "Combat/Fitness Area"),
    ]
    for pname, sport, coach_name, schedule, venue in programme_seed:
        coach_id = conn.execute("SELECT id FROM coaches WHERE name=?", (coach_name,)).fetchone()["id"]
        existing = conn.execute("SELECT id FROM programmes WHERE sport=? AND coach_id=?", (sport, coach_id)).fetchone()
        if existing:
            conn.execute("UPDATE programmes SET name=?,schedule=?,venue=?,active=1 WHERE id=?", (pname, schedule, venue, existing["id"]))
        else:
            conn.execute("INSERT INTO programmes(name,sport,coach_id,schedule,venue,active) VALUES (?,?,?,?,?,1)", (pname, sport, coach_id, schedule, venue))

    if conn.execute("SELECT COUNT(*) FROM students").fetchone()[0] == 0:
        students = [
            ("Aarav Sharma","IKC001","aarav@inspiria.edu.in","9876500001","BBA Sports Management","3rd","Football, Gym","Football trainee and fitness enthusiast."),
            ("Ananya Das","IKC002","ananya@inspiria.edu.in","9876500002","BBA Sports Management","5th","Basketball, Badminton","Basketball coaching athlete."),
            ("Rohan Singh","IKC003","rohan@inspiria.edu.in","9876500003","BBA","1st","Football","Football coaching athlete."),
            ("Meera Roy","IKC004","meera@inspiria.edu.in","9876500004","BBA Sports Management","3rd","Kick Boxing, Volleyball","Kick boxing trainee and volleyball participant."),
            ("Kabir Rai","IKC005","kabir@inspiria.edu.in","9876500005","BCA","2nd","Gym, Basketball","Gym coaching and recreational basketball."),
            ("Sneha Pradhan","IKC006","sneha@inspiria.edu.in","9876500006","BBA","4th","Basketball, Pickleball","Basketball coaching athlete."),
        ]
        conn.executemany("INSERT INTO students(name,student_id,email,phone,department,semester,sports,bio) VALUES (?,?,?,?,?,?,?,?)", students)

    # Demo student login for the prototype.
    demo_student = conn.execute("SELECT id FROM students WHERE student_id='IKC001'").fetchone()
    if demo_student:
        conn.execute("""INSERT OR IGNORE INTO users(username,password_hash,display_name,designation,role,student_id)
                        VALUES (?,?,?,?, 'student',?)""",
                     ("aarav@inspiria.edu.in", generate_password_hash("Student@2026"), "Aarav Sharma", "Student · BBA Sports Management", demo_student["id"]))

    # Only seed enrolments when there are none. These are demo records and can be replaced with real student data.
    if conn.execute("SELECT COUNT(*) FROM enrollments").fetchone()[0] == 0:
        def sid(code): return conn.execute("SELECT id FROM students WHERE student_id=?", (code,)).fetchone()["id"]
        def pid(sport): return conn.execute("SELECT id FROM programmes WHERE sport=? AND active=1 ORDER BY id DESC LIMIT 1", (sport,)).fetchone()["id"]
        conn.executemany("INSERT INTO enrollments(student_id,programme_id,status) VALUES (?,?, 'Active')", [
            (sid("IKC001"), pid("Football")), (sid("IKC002"), pid("Basketball")),
            (sid("IKC003"), pid("Football")), (sid("IKC004"), pid("Kick Boxing")),
            (sid("IKC005"), pid("Gym")), (sid("IKC006"), pid("Basketball")),
        ])

    if conn.execute("SELECT COUNT(*) FROM tournaments").fetchone()[0] == 0:
        conn.executemany("INSERT INTO tournaments(name,sport,start_date,end_date,venue,status,description) VALUES (?,?,?,?,?,?,?)", [
            ("Top Spin 4.0","Table Tennis","2026-09-25","2026-09-26","Indoor Sports Arena","Live","Inter-Departmental Table Tennis Tournament."),
            ("YouthRun 2026","Running","2026-11-22","2026-11-22","Siliguri","Upcoming","Inspiria's annual run for education."),
        ])

    if conn.execute("SELECT COUNT(*) FROM matches").fetchone()[0] == 0:
        top_spin = conn.execute("SELECT id FROM tournaments WHERE name='Top Spin 4.0'").fetchone()
        if top_spin:
            conn.executemany("""INSERT INTO matches(tournament_id,round_name,match_date,venue,team_a,team_b,score_a,score_b,status,live_note)
                              VALUES (?,?,?,?,?,?,?,?,?,?)""", [
                (top_spin["id"],"Intra / Inter Department","2026-09-25T15:00","Indoor Sports Arena","BBA SM","BBA DM","2","1","Live","Game in progress"),
                (top_spin["id"],"Inter Department","2026-09-25T15:30","Indoor Sports Arena","BBA","BCA","0","0","Scheduled",None),
                (top_spin["id"],"Inter Department","2026-09-25T16:00","Indoor Sports Arena","BBA HM","BBA E","0","0","Scheduled",None),
            ])
            live_match = conn.execute("SELECT id FROM matches WHERE tournament_id=? AND status='Live' LIMIT 1", (top_spin["id"],)).fetchone()
            if live_match:
                conn.executemany("INSERT OR IGNORE INTO match_live_stats(match_id,stat_key,stat_label,stat_value) VALUES (?,?,?,?)", [
                    (live_match["id"],"game","Current Game","4"), (live_match["id"],"games","Games Won","2 - 1"), (live_match["id"],"serve","Serving","BBA SM")
                ])


def migrate_legacy_metric_scores(conn):
    """Backfill sport-specific metric rows from V2 fixed columns when upgrading an existing database."""
    rows = conn.execute("""SELECT da.id,p.sport,da.skill_development,da.passing,da.dribbling,da.shooting,da.vision,da.fitness
                         FROM daily_assessments da JOIN enrollments e ON e.id=da.enrollment_id
                         JOIN programmes p ON p.id=e.programme_id""").fetchall()
    for row in rows:
        existing = conn.execute("SELECT 1 FROM assessment_metric_scores WHERE assessment_id=? LIMIT 1", (row["id"],)).fetchone()
        if existing:
            continue
        mapping = []
        for key, label in metrics_for_sport(row["sport"]):
            source_key = key
            if row["sport"] == "Basketball" and key == "court_vision":
                source_key = "vision"
            if source_key in row.keys() and row[source_key] is not None:
                mapping.append((key, label, row[source_key]))
        for key, label, value in mapping:
            conn.execute("INSERT OR IGNORE INTO assessment_metric_scores(assessment_id,metric_key,metric_label,score) VALUES (?,?,?,?)",
                         (row["id"], key, label, value))


def current_user():
    uid = session.get("user_id")
    if not uid:
        return None
    conn = get_db()
    user = conn.execute("SELECT * FROM users WHERE id=? AND active=1", (uid,)).fetchone()
    conn.close()
    return user


def login_required(view):
    @wraps(view)
    def wrapped(*args, **kwargs):
        if not session.get("user_id"):
            return redirect(url_for("login", next=request.path))
        return view(*args, **kwargs)
    return wrapped


def admin_required(view):
    @wraps(view)
    def wrapped(*args, **kwargs):
        user = current_user()
        if not user:
            return redirect(url_for("login", next=request.path))
        if user["role"] != "admin":
            abort(403)
        return view(*args, **kwargs)
    return wrapped


def staff_required(view):
    @wraps(view)
    def wrapped(*args, **kwargs):
        user = current_user()
        if not user:
            return redirect(url_for("login", next=request.path))
        if user["role"] not in ("admin", "coach"):
            abort(403)
        return view(*args, **kwargs)
    return wrapped


def student_required(view):
    @wraps(view)
    def wrapped(*args, **kwargs):
        user = current_user()
        if not user:
            return redirect(url_for("login", next=request.path))
        if user["role"] != "student":
            abort(403)
        return view(*args, **kwargs)
    return wrapped


def allowed_file(filename, allowed):
    return "." in filename and filename.rsplit(".", 1)[1].lower() in allowed


def save_upload(file_storage, folder, allowed):
    if not file_storage or not file_storage.filename:
        return None
    if not allowed_file(file_storage.filename, allowed):
        raise ValueError("Unsupported file type.")
    ext = file_storage.filename.rsplit(".", 1)[1].lower()
    name = f"{uuid.uuid4().hex}.{ext}"
    target_dir = os.path.join(UPLOAD_ROOT, folder)
    os.makedirs(target_dir, exist_ok=True)
    file_storage.save(os.path.join(target_dir, name))
    return f"uploads/{folder}/{name}"


def remove_static_file(relative_path):
    if not relative_path:
        return
    full = os.path.join(app.static_folder, relative_path)
    try:
        if os.path.isfile(full):
            os.remove(full)
    except OSError:
        pass


def append_student_sport(conn, student_id, sport):
    row = conn.execute("SELECT sports FROM students WHERE id=?", (student_id,)).fetchone()
    if not row:
        return
    current = [x.strip() for x in (row["sports"] or "").split(",") if x.strip()]
    if sport and sport.lower() not in {x.lower() for x in current}:
        current.append(sport)
        conn.execute("UPDATE students SET sports=? WHERE id=?", (", ".join(current), student_id))


def projection_window(period):
    today = date.today()
    if period == "weekly":
        start = today - timedelta(days=6)
    elif period == "monthly":
        start = today.replace(day=1)
    else:
        period = "yearly"
        start = today.replace(month=1, day=1)
    return period, start.isoformat(), today.isoformat()


def build_projection(conn, period):
    period, start_date, end_date = projection_window(period)
    sports = conn.execute("SELECT DISTINCT sport FROM programmes WHERE active=1 ORDER BY sport").fetchall()
    rows = []
    for sp in sports:
        sport = sp["sport"]
        attendance = conn.execute("""SELECT COUNT(a.id) sessions,COALESCE(SUM(a.present),0) present,COUNT(DISTINCT e.student_id) athletes
                                  FROM enrollments e JOIN programmes p ON p.id=e.programme_id
                                  LEFT JOIN attendance a ON a.enrollment_id=e.id AND a.session_date BETWEEN ? AND ?
                                  WHERE p.sport=? AND e.status='Active'""", (start_date,end_date,sport)).fetchone()
        perf = conn.execute("""SELECT ROUND(AVG(da.overall_performance),1) avg_performance,COUNT(da.id) rated_sessions
                            FROM daily_assessments da JOIN enrollments e ON e.id=da.enrollment_id
                            JOIN programmes p ON p.id=e.programme_id
                            WHERE p.sport=? AND da.session_date BETWEEN ? AND ?""", (sport,start_date,end_date)).fetchone()
        metrics = conn.execute("""SELECT ams.metric_key,ams.metric_label,ROUND(AVG(ams.score),1) avg_score,COUNT(*) samples
                               FROM assessment_metric_scores ams JOIN daily_assessments da ON da.id=ams.assessment_id
                               JOIN enrollments e ON e.id=da.enrollment_id JOIN programmes p ON p.id=e.programme_id
                               WHERE p.sport=? AND da.session_date BETWEEN ? AND ?
                               GROUP BY ams.metric_key,ams.metric_label ORDER BY ams.metric_label""", (sport,start_date,end_date)).fetchall()
        total = attendance["sessions"] or 0
        rate = round((attendance["present"] or 0) / total * 100, 1) if total else 0
        rows.append({"sport":sport,"athletes":attendance["athletes"] or 0,"sessions":total,"attendance_rate":rate,
                     "avg_performance":perf["avg_performance"] or 0,"rated_sessions":perf["rated_sessions"] or 0,
                     "metrics":metrics})
    return {"period":period,"from_date":start_date,"to_date":end_date,"sports":rows}


def accessible_enrollment(conn, enrollment_id, user):
    row = conn.execute("""SELECT e.*, p.name programme, p.sport, p.coach_id, p.schedule, p.venue,
                          s.name student, s.student_id student_code, s.department, s.semester,
                          c.name coach
                          FROM enrollments e
                          JOIN programmes p ON p.id=e.programme_id
                          JOIN students s ON s.id=e.student_id
                          LEFT JOIN coaches c ON c.id=p.coach_id
                          WHERE e.id=?""", (enrollment_id,)).fetchone()
    if not row:
        return None
    if user["role"] == "coach" and row["coach_id"] != user["coach_id"]:
        return None
    return row


def can_access_student(conn, student_id, user):
    if user["role"] == "admin":
        return True
    if user["role"] == "student":
        return user["student_id"] == student_id
    return conn.execute("""SELECT 1 FROM enrollments e JOIN programmes p ON p.id=e.programme_id
                           WHERE e.student_id=? AND p.coach_id=? LIMIT 1""", (student_id, user["coach_id"])).fetchone() is not None


def metrics_for_sport(sport):
    return SPORT_METRICS.get(sport, [("skill_development", "Skill Development"), ("fitness", "Fitness")])


def accessible_programme(conn, programme_id, user):
    row = conn.execute("""SELECT p.*,c.name coach FROM programmes p LEFT JOIN coaches c ON c.id=p.coach_id
                        WHERE p.id=? AND p.active=1""", (programme_id,)).fetchone()
    if not row:
        return None
    if user["role"] == "coach" and row["coach_id"] != user["coach_id"]:
        return None
    if user["role"] == "student":
        return None
    return row


def date_filters(alias="da"):
    clauses, params = [], []
    from_date = request.args.get("from", "").strip()
    to_date = request.args.get("to", "").strip()
    if from_date:
        clauses.append(f"{alias}.session_date >= ?")
        params.append(from_date)
    if to_date:
        clauses.append(f"{alias}.session_date <= ?")
        params.append(to_date)
    return clauses, params, from_date, to_date


@app.context_processor
def inject_globals():
    user = current_user(); unread = 0
    announcements = []
    if user:
        conn = get_db()
        unread = conn.execute("SELECT COUNT(*) FROM messages WHERE receiver_user_id=? AND read_at IS NULL", (user["id"],)).fetchone()[0]
        announcements = conn.execute("SELECT * FROM announcements WHERE active=1 ORDER BY created_at DESC,id DESC LIMIT 3").fetchall()
        conn.close()
    return {"now": datetime.now().strftime("%d %b %Y"), "current_user": user, "sport_metrics": SPORT_METRICS, "departments": DEPARTMENTS, "department_match": department_match, "live_score_fields": LIVE_SCORE_FIELDS, "unread_messages": unread, "global_announcements": announcements}


@app.route("/login", methods=["GET", "POST"])
def login():
    if session.get("user_id"):
        return redirect(url_for("dashboard"))
    if request.method == "POST":
        username = request.form.get("username", "").strip().lower()
        password = request.form.get("password", "")
        conn = get_db()
        user = conn.execute("SELECT * FROM users WHERE username=? AND active=1", (username,)).fetchone()
        conn.close()
        if user and check_password_hash(user["password_hash"], password):
            session.clear()
            session["user_id"] = user["id"]
            return redirect(request.args.get("next") or url_for("dashboard"))
        flash("Incorrect username or password.", "error")
    return render_template("login.html")


@app.route("/register", methods=["GET", "POST"])
def register_student():
    if session.get("user_id"):
        return redirect(url_for("dashboard"))
    if request.method == "POST":
        email = request.form.get("email", "").strip().lower()
        password = request.form.get("password", "")
        confirm = request.form.get("confirm_password", "")
        name = request.form.get("name", "").strip()
        student_code = request.form.get("student_id", "").strip() or email.split("@")[0].upper()
        department = request.form.get("department", "").strip()
        semester = request.form.get("semester", "").strip()
        phone = request.form.get("phone", "").strip()
        emergency = request.form.get("emergency_contact", "").strip()
        blood = request.form.get("blood_group", "").strip()
        medical = request.form.get("medical_history", "").strip()
        address = request.form.get("address", "").strip()
        parent_name = request.form.get("parent_name", "").strip()
        parent_phone = request.form.get("parent_phone", "").strip()
        parent_address = request.form.get("parent_address", "").strip()
        if not email.endswith("@inspiria.edu.in"):
            flash("Student registration requires an @inspiria.edu.in college email.", "error")
            return render_template("register_student.html")
        if not name or not department or not semester or not phone:
            flash("Name, department, semester and phone number are required.", "error")
            return render_template("register_student.html")
        if len(password) < 8:
            flash("Use a password with at least 8 characters.", "error")
            return render_template("register_student.html")
        if password != confirm:
            flash("Passwords do not match.", "error")
            return render_template("register_student.html")
        conn = get_db()
        try:
            if conn.execute("SELECT 1 FROM users WHERE username=?", (email,)).fetchone():
                flash("An account already exists for this college email.", "error")
                conn.close(); return render_template("register_student.html")
            student = conn.execute("SELECT * FROM students WHERE lower(email)=? OR student_id=?", (email, student_code)).fetchone()
            if student:
                conn.execute("""UPDATE students SET name=?,email=?,phone=?,department=?,semester=?,emergency_contact=?,blood_group=?,medical_history=?,
                              address=?,parent_name=?,parent_phone=?,parent_address=? WHERE id=?""",
                             (name,email,phone,department,semester,emergency,blood,medical,address,parent_name,parent_phone,parent_address,student["id"]))
                sid = student["id"]
            else:
                cur = conn.execute("""INSERT INTO students(name,student_id,email,phone,department,semester,emergency_contact,blood_group,medical_history,address,parent_name,parent_phone,parent_address)
                                  VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                                   (name,student_code,email,phone,department,semester,emergency,blood,medical,address,parent_name,parent_phone,parent_address))
                sid = cur.lastrowid
            passport_id = conn.execute("SELECT passport_id FROM students WHERE id=?", (sid,)).fetchone()["passport_id"]
            if not passport_id:
                passport_id = make_passport_id(sid)
                conn.execute("UPDATE students SET passport_id=? WHERE id=?", (passport_id, sid))
            cur = conn.execute("""INSERT INTO users(username,password_hash,display_name,designation,role,student_id) VALUES (?,?,?,?, 'student',?)""",
                               (email, generate_password_hash(password), name, f"Student · {department}", sid))
            conn.commit()
            session.clear(); session["user_id"] = cur.lastrowid
            flash("Student account created. Welcome to your Digital Sports Passport.", "success")
            return redirect(url_for("dashboard"))
        except sqlite3.IntegrityError:
            conn.rollback(); flash("This college email or student ID is already registered.", "error")
        finally:
            conn.close()
    return render_template("register_student.html")


@app.route("/logout")
def logout():
    session.clear()
    return redirect(url_for("login"))


@app.route("/")
@login_required
def dashboard():
    user = current_user()
    conn = get_db()
    today = datetime.now().strftime("%Y-%m-%d")
    if user["role"] == "student":
        student = conn.execute("SELECT * FROM students WHERE id=?", (user["student_id"],)).fetchone()
        att = conn.execute("""SELECT COUNT(*) total,COALESCE(SUM(a.present),0) present FROM attendance a
                            JOIN enrollments e ON e.id=a.enrollment_id WHERE e.student_id=?""", (user["student_id"],)).fetchone()
        attendance_rate = round(att["present"] / att["total"] * 100, 1) if att["total"] else 0
        avg_perf = conn.execute("""SELECT ROUND(AVG(da.overall_performance),1) FROM daily_assessments da
                                  JOIN enrollments e ON e.id=da.enrollment_id WHERE e.student_id=?""", (user["student_id"],)).fetchone()[0] or 0
        registrations = conn.execute("SELECT COUNT(*) FROM registrations WHERE student_id=?", (user["student_id"],)).fetchone()[0]
        coaching = conn.execute("""SELECT p.*,c.name coach,e.id enrollment_id,ROUND(AVG(da.overall_performance),1) avg_performance
                                  FROM enrollments e JOIN programmes p ON p.id=e.programme_id LEFT JOIN coaches c ON c.id=p.coach_id
                                  LEFT JOIN daily_assessments da ON da.enrollment_id=e.id WHERE e.student_id=? AND e.status='Active' GROUP BY e.id""", (user["student_id"],)).fetchall()
        recent = conn.execute("""SELECT da.*,p.sport FROM daily_assessments da JOIN enrollments e ON e.id=da.enrollment_id
                               JOIN programmes p ON p.id=e.programme_id WHERE e.student_id=? ORDER BY da.session_date DESC LIMIT 5""", (user["student_id"],)).fetchall()
        upcoming = conn.execute("SELECT * FROM tournaments WHERE start_date>=? ORDER BY start_date LIMIT 5", (today,)).fetchall()
        unread = conn.execute("SELECT COUNT(*) FROM messages WHERE receiver_user_id=? AND read_at IS NULL", (user["id"],)).fetchone()[0]
        stats = {"attendance_rate": attendance_rate, "avg_performance": avg_perf, "registrations": registrations, "unread": unread}
        conn.close()
        return render_template("dashboard.html", student=student, stats=stats, recent=recent, programmes=coaching, upcoming=upcoming, today=today)

    if user["role"] == "coach":
        scope_params = (user["coach_id"],)
        student_count = conn.execute("""SELECT COUNT(DISTINCT e.student_id) FROM enrollments e JOIN programmes p ON p.id=e.programme_id
                                        WHERE p.coach_id=? AND e.status='Active'""", scope_params).fetchone()[0]
        programme_count = conn.execute("SELECT COUNT(*) FROM programmes WHERE coach_id=? AND active=1", scope_params).fetchone()[0]
        assessments = conn.execute("SELECT COUNT(*) FROM daily_assessments WHERE coach_id=?", scope_params).fetchone()[0]
        assessed_today = conn.execute("SELECT COUNT(*) FROM daily_assessments WHERE coach_id=? AND session_date=?", (user["coach_id"], today)).fetchone()[0]
        att = conn.execute("""SELECT COUNT(*) total, COALESCE(SUM(a.present),0) present FROM attendance a JOIN enrollments e ON e.id=a.enrollment_id
                            JOIN programmes p ON p.id=e.programme_id WHERE p.coach_id=?""", scope_params).fetchone()
        recent = conn.execute("""SELECT da.*,s.name student,p.sport FROM daily_assessments da JOIN enrollments e ON e.id=da.enrollment_id
                                 JOIN students s ON s.id=e.student_id JOIN programmes p ON p.id=e.programme_id WHERE p.coach_id=?
                                 ORDER BY da.session_date DESC,da.updated_at DESC LIMIT 6""", scope_params).fetchall()
        programmes = conn.execute("""SELECT p.*,c.name coach,COUNT(DISTINCT e.id) athletes,ROUND(AVG(da.overall_performance),1) avg_performance,MAX(da.session_date) last_session
                              FROM programmes p LEFT JOIN coaches c ON c.id=p.coach_id LEFT JOIN enrollments e ON e.programme_id=p.id AND e.status='Active'
                              LEFT JOIN daily_assessments da ON da.enrollment_id=e.id WHERE p.coach_id=? AND p.active=1 GROUP BY p.id ORDER BY p.sport""", scope_params).fetchall()
    else:
        student_count = conn.execute("SELECT COUNT(*) FROM students").fetchone()[0]
        programme_count = conn.execute("SELECT COUNT(*) FROM programmes WHERE active=1").fetchone()[0]
        assessments = conn.execute("SELECT COUNT(*) FROM daily_assessments").fetchone()[0]
        assessed_today = conn.execute("SELECT COUNT(*) FROM daily_assessments WHERE session_date=?", (today,)).fetchone()[0]
        att = conn.execute("SELECT COUNT(*) total,COALESCE(SUM(present),0) present FROM attendance").fetchone()
        recent = conn.execute("""SELECT da.*,s.name student,p.sport FROM daily_assessments da JOIN enrollments e ON e.id=da.enrollment_id
                                 JOIN students s ON s.id=e.student_id JOIN programmes p ON p.id=e.programme_id ORDER BY da.session_date DESC,da.updated_at DESC LIMIT 6""").fetchall()
        programmes = conn.execute("""SELECT p.*,c.name coach,COUNT(DISTINCT e.id) athletes,ROUND(AVG(da.overall_performance),1) avg_performance,MAX(da.session_date) last_session
                              FROM programmes p LEFT JOIN coaches c ON c.id=p.coach_id LEFT JOIN enrollments e ON e.programme_id=p.id AND e.status='Active'
                              LEFT JOIN daily_assessments da ON da.enrollment_id=e.id WHERE p.active=1 GROUP BY p.id ORDER BY p.sport""").fetchall()
    attendance_rate = round((att["present"] / att["total"] * 100), 1) if att["total"] else 0
    stats = {"students":student_count,"sports":conn.execute("SELECT COUNT(*) FROM sports WHERE is_active=1").fetchone()[0],"coaching":programme_count,
             "assessments":assessments,"assessed_today":assessed_today,"attendance_rate":attendance_rate}
    sports = conn.execute("""SELECT sp.*,c.name coach FROM sports sp LEFT JOIN coaches c ON c.sport=sp.name
                             WHERE sp.is_active=1 ORDER BY CASE sp.activity_type WHEN 'Coached' THEN 0 ELSE 1 END,sp.name""").fetchall()
    upcoming = conn.execute("SELECT * FROM tournaments WHERE start_date>=? ORDER BY start_date LIMIT 4", (today,)).fetchall()
    conn.close()
    return render_template("dashboard.html", stats=stats, sports=sports[:8], recent=recent, programmes=programmes, upcoming=upcoming, today=today)

@app.route("/sports")
@login_required
def sports():
    conn = get_db()
    rows = conn.execute("""SELECT sp.*,c.name coach,
                           (SELECT COUNT(*) FROM facility_media fm WHERE fm.sport_id=sp.id) photo_count
                           FROM sports sp LEFT JOIN coaches c ON c.sport=sp.name
                           WHERE sp.is_active=1 ORDER BY sp.name""").fetchall()
    conn.close()
    return render_template("sports.html", sports=rows)


@app.route("/sport/<int:sport_id>")
@login_required
def sport_detail(sport_id):
    conn=get_db()
    sport=conn.execute("""SELECT sp.*,c.name coach FROM sports sp LEFT JOIN coaches c ON c.sport=sp.name WHERE sp.id=?""", (sport_id,)).fetchone()
    if not sport:
        conn.close(); abort(404)
    media=conn.execute("SELECT * FROM facility_media WHERE sport_id=? ORDER BY created_at DESC,id DESC", (sport_id,)).fetchall()
    conn.close()
    return render_template("sport_detail.html", sport=sport, media=media)


@app.route("/students")
@login_required
def students():
    user = current_user()
    if user["role"] == "student":
        return redirect(url_for("student_profile", student_id=user["student_id"]))
    q = request.args.get("q", "").strip()
    conn = get_db()
    like = f"%{q}%"
    base = """SELECT s.*,
              (SELECT COUNT(*) FROM enrollments e WHERE e.student_id=s.id AND e.status='Active') coaching_count,
              (SELECT COUNT(*) FROM registrations r WHERE r.student_id=s.id) tournament_count,
              (SELECT ROUND(AVG(da.overall_performance),1) FROM daily_assessments da
                 JOIN enrollments e2 ON e2.id=da.enrollment_id WHERE e2.student_id=s.id) avg_performance
              FROM students s"""
    if user["role"] == "coach":
        rows = conn.execute(base + """ WHERE EXISTS (SELECT 1 FROM enrollments ce JOIN programmes cp ON cp.id=ce.programme_id
                                   WHERE ce.student_id=s.id AND cp.coach_id=? AND ce.status='Active')
                                   AND (s.name LIKE ? OR s.student_id LIKE ? OR s.passport_id LIKE ? OR s.department LIKE ?) ORDER BY s.name""",
                            (user["coach_id"], like, like, like, like)).fetchall()
    else:
        rows = conn.execute(base + " WHERE s.name LIKE ? OR s.student_id LIKE ? OR s.passport_id LIKE ? OR s.department LIKE ? ORDER BY s.name",
                            (like, like, like, like)).fetchall()
    conn.close()
    return render_template("students.html", students=rows, q=q)


@app.route("/student/<int:student_id>")
@login_required
def student_profile(student_id):
    user = current_user()
    conn = get_db()
    student = conn.execute("SELECT * FROM students WHERE id=?", (student_id,)).fetchone()
    if not student:
        conn.close(); abort(404)
    if not can_access_student(conn, student_id, user):
        conn.close(); abort(403)
    achievements = conn.execute("SELECT * FROM achievements WHERE student_id=? ORDER BY year DESC,id DESC", (student_id,)).fetchall()
    certificates = conn.execute("SELECT * FROM certificates WHERE student_id=? ORDER BY issue_date DESC,id DESC", (student_id,)).fetchall()
    generated_certificates = conn.execute("""SELECT gc.*,t.name event_name FROM generated_certificates gc JOIN tournaments t ON t.id=gc.tournament_id
                                             WHERE gc.student_id=? ORDER BY gc.created_at DESC,gc.id DESC""", (student_id,)).fetchall()
    coaching = conn.execute("""SELECT p.*,e.id enrollment_id,e.attendance,e.total_sessions,e.status,c.name coach,
                               ROUND(AVG(da.overall_performance),1) avg_performance,MAX(da.session_date) last_session
                               FROM enrollments e JOIN programmes p ON p.id=e.programme_id
                               LEFT JOIN coaches c ON c.id=p.coach_id
                               LEFT JOIN daily_assessments da ON da.enrollment_id=e.id
                               WHERE e.student_id=? GROUP BY e.id""", (student_id,)).fetchall()
    raw_assessments = conn.execute("""SELECT da.*,p.sport,c.name coach FROM daily_assessments da
                                  JOIN enrollments e ON e.id=da.enrollment_id JOIN programmes p ON p.id=e.programme_id
                                  LEFT JOIN coaches c ON c.id=da.coach_id
                                  WHERE e.student_id=? ORDER BY da.session_date DESC,da.id DESC LIMIT 12""", (student_id,)).fetchall()
    assessments = []
    for row in raw_assessments:
        item = dict(row)
        item["metrics"] = conn.execute("SELECT metric_label,score FROM assessment_metric_scores WHERE assessment_id=? ORDER BY id", (row["id"],)).fetchall()
        assessments.append(item)
    tournaments = conn.execute("""SELECT t.*,r.team_name FROM registrations r JOIN tournaments t ON t.id=r.tournament_id
                                  WHERE r.student_id=? ORDER BY t.start_date DESC""", (student_id,)).fetchall()
    metric_summary = conn.execute("""SELECT ams.metric_label,ROUND(AVG(ams.score),1) avg_score,COUNT(*) samples
                          FROM assessment_metric_scores ams JOIN daily_assessments da ON da.id=ams.assessment_id
                          JOIN enrollments e ON e.id=da.enrollment_id WHERE e.student_id=?
                          GROUP BY ams.metric_key,ams.metric_label ORDER BY avg_score DESC,ams.metric_label LIMIT 8""", (student_id,)).fetchall()
    attendance = conn.execute("""SELECT COUNT(*) total,COALESCE(SUM(a.present),0) present FROM attendance a
                                JOIN enrollments e ON e.id=a.enrollment_id WHERE e.student_id=?""", (student_id,)).fetchone()
    attendance_rate = round(attendance["present"] / attendance["total"] * 100, 1) if attendance["total"] else 0
    conn.close()
    return render_template("student_profile.html", student=student, achievements=achievements, certificates=certificates,
                           coaching=coaching, assessments=assessments, tournaments=tournaments, metric_summary=metric_summary,
                           generated_certificates=generated_certificates, attendance_rate=attendance_rate)


@app.route("/admin/student/<int:student_id>/photo", methods=["POST"])
@admin_required
def upload_student_photo(student_id):
    upload = request.files.get("photo")
    try:
        path = save_upload(upload, "profiles/students", ALLOWED_IMAGE_EXTENSIONS)
        if not path: raise ValueError("Choose a student profile photo.")
        conn=get_db(); row=conn.execute("SELECT photo_path FROM students WHERE id=?",(student_id,)).fetchone()
        if not row: conn.close(); abort(404)
        old=row["photo_path"]; conn.execute("UPDATE students SET photo_path=? WHERE id=?",(path,student_id)); conn.commit(); conn.close()
        remove_static_file(old); flash("Student profile photo updated.","success")
    except ValueError as exc: flash(str(exc),"error")
    return redirect(request.referrer or url_for("student_profile",student_id=student_id))


@app.route("/admin/student/<int:student_id>/photo/delete", methods=["POST"])
@admin_required
def delete_student_photo(student_id):
    conn=get_db(); row=conn.execute("SELECT photo_path FROM students WHERE id=?",(student_id,)).fetchone()
    if row:
        old=row["photo_path"]; conn.execute("UPDATE students SET photo_path=NULL WHERE id=?",(student_id,)); conn.commit(); remove_static_file(old)
    conn.close(); flash("Student profile photo removed.","success")
    return redirect(request.referrer or url_for("student_profile",student_id=student_id))


@app.route("/admin/coach/<int:coach_id>/photo", methods=["POST"])
@admin_required
def upload_coach_photo(coach_id):
    upload=request.files.get("photo")
    try:
        path=save_upload(upload,"profiles/coaches",ALLOWED_IMAGE_EXTENSIONS)
        if not path: raise ValueError("Choose a coach profile photo.")
        conn=get_db(); row=conn.execute("SELECT photo_path FROM coaches WHERE id=?",(coach_id,)).fetchone()
        if not row: conn.close(); abort(404)
        old=row["photo_path"]; conn.execute("UPDATE coaches SET photo_path=? WHERE id=?",(path,coach_id)); conn.commit(); conn.close()
        remove_static_file(old); flash("Coach profile photo updated.","success")
    except ValueError as exc: flash(str(exc),"error")
    return redirect(request.referrer or url_for("admin"))


@app.route("/admin/coach/<int:coach_id>/photo/delete", methods=["POST"])
@admin_required
def delete_coach_photo(coach_id):
    conn=get_db(); row=conn.execute("SELECT photo_path FROM coaches WHERE id=?",(coach_id,)).fetchone()
    if row:
        old=row["photo_path"]; conn.execute("UPDATE coaches SET photo_path=NULL WHERE id=?",(coach_id,)); conn.commit(); remove_static_file(old)
    conn.close(); flash("Coach profile photo removed.","success")
    return redirect(request.referrer or url_for("admin"))


@app.route("/students/add", methods=["POST"])
@admin_required
def add_student():
    fields = ["name","student_id","email","phone","department","semester","sports","bio","emergency_contact","blood_group","medical_history","address","parent_name","parent_phone","parent_address"]
    data = [request.form.get(k, "").strip() for k in fields]
    conn = get_db()
    try:
        cur=conn.execute("""INSERT INTO students(name,student_id,email,phone,department,semester,sports,bio,emergency_contact,blood_group,medical_history,address,parent_name,parent_phone,parent_address)
                      VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""", data)
        conn.execute("UPDATE students SET passport_id=? WHERE id=?", (make_passport_id(cur.lastrowid),cur.lastrowid))
        conn.commit(); flash("Student profile created with a unique Sports Passport ID.", "success")
    except sqlite3.IntegrityError:
        flash("Student ID already exists.", "error")
    finally:
        conn.close()
    return redirect(url_for("students"))


@app.route("/student/<int:student_id>/update", methods=["POST"])
@admin_required
def update_student(student_id):
    fields = ["name","student_id","email","phone","department","semester","sports","bio","emergency_contact","blood_group","medical_history","address","parent_name","parent_phone","parent_address"]
    vals = [request.form.get(k, "").strip() for k in fields]
    conn = get_db()
    try:
        conn.execute("""UPDATE students SET name=?,student_id=?,email=?,phone=?,department=?,semester=?,sports=?,bio=?,emergency_contact=?,blood_group=?,medical_history=?,address=?,parent_name=?,parent_phone=?,parent_address=?
                        WHERE id=?""", (*vals, student_id))
        conn.commit(); flash("Student profile updated.", "success")
    except sqlite3.IntegrityError:
        flash("That Student ID is already assigned to another profile.", "error")
    finally:
        conn.close()
    return redirect(url_for("student_profile", student_id=student_id))


@app.route("/students/import", methods=["POST"])
@admin_required
def import_students():
    upload = request.files.get("file")
    if not upload or not upload.filename:
        flash("Choose a CSV file to import.", "error")
        return redirect(url_for("students"))
    try:
        text_data = upload.stream.read().decode("utf-8-sig")
        reader = csv.DictReader(io.StringIO(text_data))
    except Exception:
        flash("The CSV could not be read. Save it as UTF-8 CSV and try again.", "error")
        return redirect(url_for("students"))
    required = {"name", "student_id"}
    if not reader.fieldnames or not required.issubset(set(reader.fieldnames)):
        flash("CSV must contain at least name and student_id columns.", "error")
        return redirect(url_for("students"))
    conn = get_db(); added = 0; updated = 0; skipped = 0
    for row in reader:
        name = (row.get("name") or "").strip(); code = (row.get("student_id") or "").strip()
        if not name or not code:
            skipped += 1; continue
        extra_fields = ["email","phone","department","semester","sports","bio","emergency_contact","blood_group","medical_history","address","parent_name","parent_phone","parent_address"]
        vals = [(row.get(k) or "").strip() for k in extra_fields]
        existing = conn.execute("SELECT id FROM students WHERE student_id=?", (code,)).fetchone()
        if existing:
            conn.execute("""UPDATE students SET name=?,email=?,phone=?,department=?,semester=?,sports=?,bio=?,emergency_contact=?,blood_group=?,medical_history=?,address=?,parent_name=?,parent_phone=?,parent_address=? WHERE id=?""",
                         (name, *vals, existing["id"]))
            updated += 1
        else:
            cur=conn.execute("""INSERT INTO students(name,student_id,email,phone,department,semester,sports,bio,emergency_contact,blood_group,medical_history,address,parent_name,parent_phone,parent_address)
                            VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""", (name, code, *vals))
            conn.execute("UPDATE students SET passport_id=? WHERE id=?", (make_passport_id(cur.lastrowid),cur.lastrowid))
            added += 1
    conn.commit(); conn.close()
    flash(f"Student import complete: {added} added, {updated} updated, {skipped} skipped.", "success")
    return redirect(url_for("students"))


@app.route("/student/<int:student_id>/achievement", methods=["POST"])
@admin_required
def add_achievement(student_id):
    vals = [request.form.get(k, "").strip() for k in ["title","sport","level","year","description"]]
    if not vals[0]:
        flash("Achievement title is required.", "error")
    else:
        conn = get_db(); conn.execute("INSERT INTO achievements(student_id,title,sport,level,year,description) VALUES (?,?,?,?,?,?)", (student_id, *vals)); conn.commit(); conn.close()
        flash("Achievement added to the Sports Passport.", "success")
    return redirect(url_for("student_profile", student_id=student_id))


@app.route("/student/<int:student_id>/certificate", methods=["POST"])
@admin_required
def add_certificate(student_id):
    vals = [request.form.get(k, "").strip() for k in ["title","issuer","issue_date","certificate_no"]]
    if not vals[0]:
        flash("Certificate title is required.", "error")
    else:
        conn = get_db(); conn.execute("INSERT INTO certificates(student_id,title,issuer,issue_date,certificate_no) VALUES (?,?,?,?,?)", (student_id, *vals)); conn.commit(); conn.close()
        flash("Certificate added to the Sports Passport.", "success")
    return redirect(url_for("student_profile", student_id=student_id))


@app.route("/admin/students/bulk-sport", methods=["POST"])
@admin_required
def bulk_assign_sport():
    sport=request.form.get("sport","").strip(); student_ids=request.form.getlist("student_ids")
    if not sport or not student_ids:
        flash("Choose a sport/activity and at least one student.","error"); return redirect(request.referrer or url_for("admin"))
    conn=get_db(); valid=conn.execute("SELECT 1 FROM sports WHERE name=? AND is_active=1",(sport,)).fetchone()
    if not valid:
        conn.close(); flash("Active sport/activity not found.","error"); return redirect(url_for("admin"))
    updated=0
    for raw in student_ids:
        if str(raw).isdigit() and conn.execute("SELECT 1 FROM students WHERE id=?",(int(raw),)).fetchone():
            before=conn.execute("SELECT sports FROM students WHERE id=?",(int(raw),)).fetchone()["sports"] or ""
            append_student_sport(conn,int(raw),sport)
            after=conn.execute("SELECT sports FROM students WHERE id=?",(int(raw),)).fetchone()["sports"] or ""
            if before != after: updated+=1
    conn.commit(); conn.close(); flash(f"{sport} added to {updated} student Sports Passport profile(s).","success")
    return redirect(request.referrer or url_for("admin"))


@app.route("/admin/enrollments/bulk", methods=["POST"])
@admin_required
def bulk_enroll_students():
    programme_id=request.form.get("programme_id",type=int); student_ids=request.form.getlist("student_ids")
    if not programme_id or not student_ids:
        flash("Choose a coaching programme and at least one student.","error"); return redirect(request.referrer or url_for("admin"))
    conn=get_db(); programme=conn.execute("SELECT * FROM programmes WHERE id=? AND active=1",(programme_id,)).fetchone()
    if not programme: conn.close(); flash("Coaching programme not found.","error"); return redirect(url_for("admin"))
    added=reactivated=skipped=0
    for raw in student_ids:
        if not str(raw).isdigit(): continue
        sid=int(raw); student=conn.execute("SELECT id FROM students WHERE id=?",(sid,)).fetchone()
        if not student: continue
        existing=conn.execute("SELECT id,status FROM enrollments WHERE student_id=? AND programme_id=?",(sid,programme_id)).fetchone()
        if existing:
            if existing["status"] != "Active":
                conn.execute("UPDATE enrollments SET status='Active' WHERE id=?",(existing["id"],)); reactivated+=1
            else: skipped+=1
        else:
            conn.execute("INSERT INTO enrollments(student_id,programme_id,status) VALUES (?,?,'Active')",(sid,programme_id)); added+=1
        append_student_sport(conn,sid,programme["sport"])
    conn.commit(); conn.close(); flash(f"Bulk assignment complete: {added} enrolled, {reactivated} reactivated, {skipped} already active.","success")
    return redirect(request.referrer or url_for("admin"))


@app.route("/coaching/schedule")
@login_required
def coaching_schedule():
    user = current_user()
    month = request.args.get("month", "").strip()
    conn = get_db()
    programmes = conn.execute("""SELECT p.*,c.name coach FROM programmes p
                               LEFT JOIN coaches c ON c.id=p.coach_id
                               WHERE p.active=1 ORDER BY p.sport,p.name""").fetchall()
    clauses = []
    params = []
    if month:
        try:
            datetime.strptime(month, "%Y-%m")
            clauses.append("substr(cs.session_date,1,7)=?")
            params.append(month)
        except ValueError:
            month = ""
    sql = """SELECT cs.*,p.name programme,p.sport,p.coach_id,p.schedule recurring_schedule,
                    c.name coach
             FROM coaching_sessions cs
             JOIN programmes p ON p.id=cs.programme_id
             LEFT JOIN coaches c ON c.id=p.coach_id"""
    if clauses:
        sql += " WHERE " + " AND ".join(clauses)
    sql += " ORDER BY cs.session_date ASC,cs.start_time ASC,cs.id ASC"
    sessions = conn.execute(sql, params).fetchall()
    editable_programme_ids = []
    if user["role"] == "admin":
        editable_programme_ids = [p["id"] for p in programmes]
    elif user["role"] == "coach":
        editable_programme_ids = [p["id"] for p in programmes if p["coach_id"] == user["coach_id"]]
    conn.close()
    return render_template("coaching_schedule.html", programmes=programmes, sessions=sessions, month=month,
                           editable_programme_ids=editable_programme_ids, today=date.today().isoformat())


@app.route("/coaching/schedule/add", methods=["POST"])
@staff_required
def add_coaching_session():
    user = current_user()
    programme_id = request.form.get("programme_id", type=int)
    session_date = request.form.get("session_date", "").strip()
    start_time = request.form.get("start_time", "").strip()
    end_time = request.form.get("end_time", "").strip()
    venue = request.form.get("venue", "").strip()
    notes = request.form.get("notes", "").strip()
    status = request.form.get("status", "Scheduled").strip() or "Scheduled"
    conn = get_db()
    if not programme_id or not can_edit_programme(conn, programme_id, user):
        conn.close(); abort(403)
    if not session_date or not start_time or not end_time:
        conn.close(); flash("Session date, start time and end time are required.", "error"); return redirect(url_for("coaching_schedule"))
    if end_time <= start_time:
        conn.close(); flash("End time must be later than start time.", "error"); return redirect(url_for("coaching_schedule"))
    conn.execute("""INSERT INTO coaching_sessions(programme_id,session_date,start_time,end_time,venue,notes,status,created_by,updated_by)
                    VALUES (?,?,?,?,?,?,?,?,?)""",
                 (programme_id,session_date,start_time,end_time,venue,notes,status,user["id"],user["id"]))
    conn.commit(); conn.close()
    flash("Coaching session added to the schedule.", "success")
    return redirect(url_for("coaching_schedule", month=session_date[:7]))


@app.route("/coaching/schedule/<int:session_id>/edit", methods=["POST"])
@staff_required
def edit_coaching_session(session_id):
    user = current_user(); conn = get_db()
    row = conn.execute("SELECT * FROM coaching_sessions WHERE id=?", (session_id,)).fetchone()
    if not row:
        conn.close(); abort(404)
    if not can_edit_programme(conn, row["programme_id"], user):
        conn.close(); abort(403)
    session_date = request.form.get("session_date", row["session_date"]).strip()
    start_time = request.form.get("start_time", row["start_time"]).strip()
    end_time = request.form.get("end_time", row["end_time"]).strip()
    venue = request.form.get("venue", row["venue"] or "").strip()
    notes = request.form.get("notes", row["notes"] or "").strip()
    status = request.form.get("status", row["status"] or "Scheduled").strip() or "Scheduled"
    if not session_date or not start_time or not end_time:
        conn.close(); flash("Session date, start time and end time are required.", "error"); return redirect(url_for("coaching_schedule"))
    if end_time <= start_time:
        conn.close(); flash("End time must be later than start time.", "error"); return redirect(url_for("coaching_schedule", month=session_date[:7]))
    conn.execute("""UPDATE coaching_sessions SET session_date=?,start_time=?,end_time=?,venue=?,notes=?,status=?,
                    updated_by=?,updated_at=CURRENT_TIMESTAMP WHERE id=?""",
                 (session_date,start_time,end_time,venue,notes,status,user["id"],session_id))
    conn.commit(); conn.close()
    flash("Coaching session updated.", "success")
    return redirect(url_for("coaching_schedule", month=session_date[:7]))


@app.route("/coaching/schedule/<int:session_id>/delete", methods=["POST"])
@staff_required
def delete_coaching_session(session_id):
    user = current_user(); conn = get_db()
    row = conn.execute("SELECT programme_id,session_date FROM coaching_sessions WHERE id=?", (session_id,)).fetchone()
    if not row:
        conn.close(); abort(404)
    if not can_edit_programme(conn, row["programme_id"], user):
        conn.close(); abort(403)
    month = (row["session_date"] or "")[:7]
    conn.execute("DELETE FROM coaching_sessions WHERE id=?", (session_id,))
    conn.commit(); conn.close()
    flash("Coaching session removed from the schedule.", "success")
    return redirect(url_for("coaching_schedule", month=month))


@app.route("/coaching")
@staff_required
def coaching():
    user = current_user()
    conn = get_db()
    if user["role"] == "coach":
        programmes = conn.execute("""SELECT p.*,c.name coach,COUNT(e.id) enrolled FROM programmes p
                                   LEFT JOIN coaches c ON c.id=p.coach_id LEFT JOIN enrollments e ON e.programme_id=p.id AND e.status='Active'
                                   WHERE p.coach_id=? AND p.active=1 GROUP BY p.id ORDER BY p.sport""", (user["coach_id"],)).fetchall()
        enrollments = conn.execute("""SELECT e.*,s.name student,s.student_id,p.name programme,p.sport,p.coach_id,
                                      ROUND(AVG(da.overall_performance),1) avg_performance,MAX(da.session_date) last_session
                                      FROM enrollments e JOIN students s ON s.id=e.student_id JOIN programmes p ON p.id=e.programme_id
                                      LEFT JOIN daily_assessments da ON da.enrollment_id=e.id
                                      WHERE p.coach_id=? AND e.status='Active' GROUP BY e.id ORDER BY s.name""", (user["coach_id"],)).fetchall()
    else:
        programmes = conn.execute("""SELECT p.*,c.name coach,COUNT(e.id) enrolled FROM programmes p
                                   LEFT JOIN coaches c ON c.id=p.coach_id LEFT JOIN enrollments e ON e.programme_id=p.id AND e.status='Active'
                                   WHERE p.active=1 GROUP BY p.id ORDER BY p.sport""").fetchall()
        enrollments = conn.execute("""SELECT e.*,s.name student,s.student_id,p.name programme,p.sport,p.coach_id,
                                      ROUND(AVG(da.overall_performance),1) avg_performance,MAX(da.session_date) last_session
                                      FROM enrollments e JOIN students s ON s.id=e.student_id JOIN programmes p ON p.id=e.programme_id
                                      LEFT JOIN daily_assessments da ON da.enrollment_id=e.id
                                      WHERE e.status='Active' GROUP BY e.id ORDER BY p.sport,s.name""").fetchall()
    conn.close()
    return render_template("coaching.html", programmes=programmes, enrollments=enrollments)


@app.route("/coaching/athlete/<int:enrollment_id>")
@staff_required
def coaching_athlete(enrollment_id):
    user = current_user()
    conn = get_db()
    enrollment = accessible_enrollment(conn, enrollment_id, user)
    if not enrollment:
        conn.close(); abort(403)
    raw_history = conn.execute("""SELECT da.*,c.name coach FROM daily_assessments da LEFT JOIN coaches c ON c.id=da.coach_id
                             WHERE da.enrollment_id=? ORDER BY da.session_date DESC,da.id DESC LIMIT 30""", (enrollment_id,)).fetchall()
    history = []
    for row in raw_history:
        item = dict(row)
        item["metrics"] = conn.execute("SELECT metric_key,metric_label,score FROM assessment_metric_scores WHERE assessment_id=? ORDER BY id", (row["id"],)).fetchall()
        history.append(item)
    metric_rows = conn.execute("""SELECT ams.metric_key,ams.metric_label,ROUND(AVG(ams.score),1) avg_score,COUNT(*) samples
                                 FROM assessment_metric_scores ams JOIN daily_assessments da ON da.id=ams.assessment_id
                                 WHERE da.enrollment_id=? GROUP BY ams.metric_key,ams.metric_label""", (enrollment_id,)).fetchall()
    metric_map = {r["metric_key"]: r for r in metric_rows}
    metrics = metrics_for_sport(enrollment["sport"])
    metric_averages = [{"key": key, "label": label, "avg_score": metric_map.get(key)["avg_score"] if metric_map.get(key) else None,
                        "samples": metric_map.get(key)["samples"] if metric_map.get(key) else 0} for key, label in metrics]
    common_averages = conn.execute("""SELECT ROUND(AVG(attendance_score),1) attendance_score,
                              ROUND(AVG(overall_performance),1) overall_performance
                              FROM daily_assessments WHERE enrollment_id=?""", (enrollment_id,)).fetchone()
    attendance = conn.execute("SELECT COUNT(*) total,COALESCE(SUM(present),0) present FROM attendance WHERE enrollment_id=?", (enrollment_id,)).fetchone()
    attendance_rate = round(attendance["present"] / attendance["total"] * 100, 1) if attendance["total"] else 0
    conn.close()
    return render_template("coaching_athlete.html", enrollment=enrollment, history=history, metric_averages=metric_averages,
                           common_averages=common_averages, metrics=metrics, attendance_rate=attendance_rate,
                           today=datetime.now().strftime("%Y-%m-%d"))


@app.route("/coaching/athlete/<int:enrollment_id>/assessment", methods=["POST"])
@staff_required
def save_assessment(enrollment_id):
    user = current_user()
    conn = get_db()
    enrollment = accessible_enrollment(conn, enrollment_id, user)
    if not enrollment:
        conn.close(); abort(403)

    def score_value(name, allow_blank=True):
        raw = request.form.get(name, "").strip()
        if raw == "" and allow_blank:
            return None
        try:
            value = int(raw)
        except ValueError:
            raise ValueError("All ratings must be numbers from 1 to 10.")
        if value < 1 or value > 10:
            raise ValueError("All ratings must be between 1 and 10.")
        return value

    try:
        attendance_score = score_value("attendance_score")
        overall = score_value("overall_performance")
        metric_values = []
        for key, label in metrics_for_sport(enrollment["sport"]):
            value = score_value(f"metric_{key}")
            if value is not None:
                metric_values.append((key, label, value))
    except ValueError as exc:
        conn.close(); flash(str(exc), "error")
        return redirect(url_for("coaching_athlete", enrollment_id=enrollment_id))

    date = request.form.get("session_date") or datetime.now().strftime("%Y-%m-%d")
    attendance_status = request.form.get("attendance_status", "Present")
    remarks = request.form.get("remarks", "").strip()
    coach_id = enrollment["coach_id"]
    if attendance_score is None:
        attendance_score = ATTENDANCE_SCORES.get(attendance_status)

    legacy = {"skill_development": None, "passing": None, "dribbling": None, "shooting": None, "vision": None, "fitness": None}
    for key, _label, value in metric_values:
        if key in legacy:
            legacy[key] = value

    conn.execute("""INSERT INTO daily_assessments(
                    enrollment_id,coach_id,session_date,attendance_status,attendance_score,skill_development,
                    passing,dribbling,shooting,vision,fitness,overall_performance,remarks,updated_at)
                    VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,CURRENT_TIMESTAMP)
                    ON CONFLICT(enrollment_id,session_date) DO UPDATE SET
                    coach_id=excluded.coach_id,attendance_status=excluded.attendance_status,
                    attendance_score=excluded.attendance_score,skill_development=excluded.skill_development,
                    passing=excluded.passing,dribbling=excluded.dribbling,shooting=excluded.shooting,
                    vision=excluded.vision,fitness=excluded.fitness,overall_performance=excluded.overall_performance,
                    remarks=excluded.remarks,updated_at=CURRENT_TIMESTAMP""",
                 (enrollment_id, coach_id, date, attendance_status, attendance_score, legacy["skill_development"],
                  legacy["passing"], legacy["dribbling"], legacy["shooting"], legacy["vision"], legacy["fitness"],
                  overall, remarks))
    assessment_id = conn.execute("SELECT id FROM daily_assessments WHERE enrollment_id=? AND session_date=?", (enrollment_id, date)).fetchone()["id"]
    conn.execute("DELETE FROM assessment_metric_scores WHERE assessment_id=?", (assessment_id,))
    for key, label, value in metric_values:
        conn.execute("INSERT INTO assessment_metric_scores(assessment_id,metric_key,metric_label,score) VALUES (?,?,?,?)",
                     (assessment_id, key, label, value))
    present = 0 if attendance_status == "Absent" else 1
    conn.execute("INSERT OR REPLACE INTO attendance(enrollment_id,session_date,present) VALUES (?,?,?)", (enrollment_id, date, present))
    totals = conn.execute("SELECT COUNT(*) total, COALESCE(SUM(present),0) present FROM attendance WHERE enrollment_id=?", (enrollment_id,)).fetchone()
    conn.execute("UPDATE enrollments SET attendance=?,total_sessions=? WHERE id=?", (totals["present"], totals["total"], enrollment_id))
    conn.commit(); conn.close()
    flash("Daily coaching assessment saved.", "success")
    return redirect(url_for("coaching_athlete", enrollment_id=enrollment_id))


@app.route("/coaching/enroll", methods=["POST"])
@admin_required
def enroll():
    conn = get_db()
    student_code = request.form.get("student_code", "").strip()
    programme_id = request.form.get("programme_id", "").strip()
    student = conn.execute("SELECT id FROM students WHERE student_id=?", (student_code,)).fetchone()
    if not student:
        conn.close(); flash("Student ID not found.", "error"); return redirect(url_for("coaching"))
    try:
        conn.execute("INSERT INTO enrollments(student_id,programme_id,status) VALUES (?,?,'Active')", (student["id"], programme_id))
        conn.commit(); flash("Student enrolled in coaching.", "success")
    except sqlite3.IntegrityError:
        flash("Student is already enrolled in that programme.", "error")
    finally:
        conn.close()
    return redirect(url_for("coaching"))


@app.route("/coaching/programme/<int:programme_id>/attendance", methods=["GET", "POST"])
@staff_required
def programme_attendance(programme_id):
    user = current_user(); conn = get_db()
    programme = accessible_programme(conn, programme_id, user)
    if not programme:
        conn.close(); abort(403)
    session_date = (request.form.get("date") if request.method == "POST" else request.args.get("date")) or datetime.now().strftime("%Y-%m-%d")
    athletes = conn.execute("""SELECT e.id enrollment_id,s.id student_pk,s.name,s.student_id,s.department,s.semester,
                              da.attendance_status,da.attendance_score,da.overall_performance
                              FROM enrollments e JOIN students s ON s.id=e.student_id
                              LEFT JOIN daily_assessments da ON da.enrollment_id=e.id AND da.session_date=?
                              WHERE e.programme_id=? AND e.status='Active' ORDER BY s.name""", (session_date, programme_id)).fetchall()
    if request.method == "POST":
        valid_ids = {a["enrollment_id"] for a in athletes}
        for enrollment_id in valid_ids:
            status = request.form.get(f"status_{enrollment_id}", "Present")
            if status not in ATTENDANCE_SCORES:
                status = "Present"
            score = ATTENDANCE_SCORES[status]
            present = 0 if status == "Absent" else 1
            conn.execute("INSERT OR REPLACE INTO attendance(enrollment_id,session_date,present) VALUES (?,?,?)", (enrollment_id, session_date, present))
            conn.execute("""INSERT INTO daily_assessments(enrollment_id,coach_id,session_date,attendance_status,attendance_score,updated_at)
                            VALUES (?,?,?,?,?,CURRENT_TIMESTAMP)
                            ON CONFLICT(enrollment_id,session_date) DO UPDATE SET
                            coach_id=excluded.coach_id,attendance_status=excluded.attendance_status,
                            attendance_score=excluded.attendance_score,updated_at=CURRENT_TIMESTAMP""",
                         (enrollment_id, programme["coach_id"], session_date, status, score))
            totals = conn.execute("SELECT COUNT(*) total,COALESCE(SUM(present),0) present FROM attendance WHERE enrollment_id=?", (enrollment_id,)).fetchone()
            conn.execute("UPDATE enrollments SET attendance=?,total_sessions=? WHERE id=?", (totals["present"], totals["total"], enrollment_id))
        conn.commit(); conn.close()
        flash(f"Attendance saved for {len(valid_ids)} athlete(s) on {session_date}.", "success")
        return redirect(url_for("programme_attendance", programme_id=programme_id, date=session_date))
    conn.close()
    return render_template("programme_attendance.html", programme=programme, athletes=athletes, session_date=session_date)


@app.route("/reports")
@login_required
def reports():
    user = current_user(); conn = get_db()
    if user["role"] == "coach":
        programmes = conn.execute("SELECT p.*,c.name coach FROM programmes p LEFT JOIN coaches c ON c.id=p.coach_id WHERE p.active=1 AND p.coach_id=? ORDER BY p.sport", (user["coach_id"],)).fetchall()
    elif user["role"] == "student":
        programmes = conn.execute("""SELECT DISTINCT p.*,c.name coach FROM programmes p
                                   JOIN enrollments e ON e.programme_id=p.id LEFT JOIN coaches c ON c.id=p.coach_id
                                   WHERE p.active=1 AND e.status='Active' AND e.student_id=? ORDER BY p.sport""", (user["student_id"],)).fetchall()
    else:
        programmes = conn.execute("SELECT p.*,c.name coach FROM programmes p LEFT JOIN coaches c ON c.id=p.coach_id WHERE p.active=1 ORDER BY p.sport").fetchall()
    allowed = {p["id"] for p in programmes}
    raw_pid = request.args.get("programme_id", "").strip()
    programme_id = int(raw_pid) if raw_pid.isdigit() and int(raw_pid) in allowed else None
    from_date = request.args.get("from", "").strip(); to_date = request.args.get("to", "").strip()

    scope_sql = ""; scope_params = []
    if user["role"] == "coach":
        scope_sql += " AND p.coach_id=?"; scope_params.append(user["coach_id"])
    elif user["role"] == "student":
        scope_sql += " AND s.id=?"; scope_params.append(user["student_id"])
    if programme_id:
        scope_sql += " AND p.id=?"; scope_params.append(programme_id)

    enrollments = conn.execute("""SELECT e.id,s.name,s.student_id,s.department,p.id programme_id,p.name programme,p.sport
                                FROM enrollments e JOIN students s ON s.id=e.student_id JOIN programmes p ON p.id=e.programme_id
                                WHERE e.status='Active'""" + scope_sql + " ORDER BY p.sport,s.name", scope_params).fetchall()
    athlete_rows = []
    for e in enrollments:
        date_sql = ""; date_params = []
        if from_date:
            date_sql += " AND session_date>=?"; date_params.append(from_date)
        if to_date:
            date_sql += " AND session_date<=?"; date_params.append(to_date)
        att = conn.execute("SELECT COUNT(*) total,COALESCE(SUM(present),0) present FROM attendance WHERE enrollment_id=?" + date_sql, (e["id"], *date_params)).fetchone()
        perf = conn.execute("SELECT ROUND(AVG(overall_performance),1) avg_perf,COUNT(overall_performance) ratings,MAX(session_date) last_session FROM daily_assessments WHERE enrollment_id=?" + date_sql, (e["id"], *date_params)).fetchone()
        rate = round(att["present"] / att["total"] * 100, 1) if att["total"] else 0
        athlete_rows.append({**dict(e), "sessions": att["total"], "present": att["present"], "attendance_rate": rate,
                             "avg_performance": perf["avg_perf"], "ratings": perf["ratings"], "last_session": perf["last_session"]})

    total_sessions = sum(r["sessions"] for r in athlete_rows); total_present = sum(r["present"] for r in athlete_rows)
    overall_attendance = round(total_present / total_sessions * 100, 1) if total_sessions else 0
    perf_values = [r["avg_performance"] for r in athlete_rows if r["avg_performance"] is not None]
    overall_performance = round(sum(perf_values)/len(perf_values), 1) if perf_values else 0

    metric_scope = " WHERE 1=1"; metric_params = []
    if user["role"] == "coach":
        metric_scope += " AND p.coach_id=?"; metric_params.append(user["coach_id"])
    elif user["role"] == "student":
        metric_scope += " AND e.student_id=?"; metric_params.append(user["student_id"])
    if programme_id:
        metric_scope += " AND p.id=?"; metric_params.append(programme_id)
    if from_date:
        metric_scope += " AND da.session_date>=?"; metric_params.append(from_date)
    if to_date:
        metric_scope += " AND da.session_date<=?"; metric_params.append(to_date)
    metric_summary = conn.execute("""SELECT ams.metric_label,ROUND(AVG(ams.score),1) avg_score,COUNT(*) samples
                          FROM assessment_metric_scores ams JOIN daily_assessments da ON da.id=ams.assessment_id
                          JOIN enrollments e ON e.id=da.enrollment_id JOIN programmes p ON p.id=e.programme_id""" + metric_scope +
                          " GROUP BY ams.metric_key,ams.metric_label ORDER BY avg_score DESC,ams.metric_label", metric_params).fetchall()
    trend = conn.execute("""SELECT da.session_date,ROUND(AVG(da.overall_performance),1) avg_performance,COUNT(DISTINCT da.enrollment_id) athletes
                         FROM daily_assessments da JOIN enrollments e ON e.id=da.enrollment_id JOIN programmes p ON p.id=e.programme_id""" +
                         metric_scope + " GROUP BY da.session_date ORDER BY da.session_date DESC LIMIT 10", metric_params).fetchall()
    trend = list(reversed(trend))
    projections = {period: build_projection(conn, period) for period in ("weekly","monthly","yearly")}
    conn.close()
    summary = {"athletes": len(athlete_rows), "sessions": total_sessions, "attendance_rate": overall_attendance, "avg_performance": overall_performance}
    return render_template("reports.html", programmes=programmes, programme_id=programme_id, from_date=from_date, to_date=to_date,
                           athlete_rows=athlete_rows, metric_summary=metric_summary, trend=trend, summary=summary, projections=projections)


@app.route("/reports/export.csv")
@login_required
def export_reports():
    user = current_user(); conn = get_db()
    programme_id = request.args.get("programme_id", "").strip(); from_date = request.args.get("from", "").strip(); to_date = request.args.get("to", "").strip()
    sql = """SELECT da.session_date,s.student_id,s.name,p.sport,p.name programme,da.attendance_status,da.attendance_score,
             da.overall_performance,da.remarks,c.name coach FROM daily_assessments da
             JOIN enrollments e ON e.id=da.enrollment_id JOIN students s ON s.id=e.student_id
             JOIN programmes p ON p.id=e.programme_id LEFT JOIN coaches c ON c.id=da.coach_id WHERE 1=1"""
    params = []
    if user["role"] == "coach":
        sql += " AND p.coach_id=?"; params.append(user["coach_id"])
    elif user["role"] == "student":
        sql += " AND s.id=?"; params.append(user["student_id"])
    if programme_id.isdigit():
        pid = int(programme_id)
        if user["role"] == "coach":
            prog = accessible_programme(conn, pid, user)
            if not prog: conn.close(); abort(403)
        elif user["role"] == "student":
            ok = conn.execute("SELECT 1 FROM enrollments WHERE student_id=? AND programme_id=? AND status='Active'", (user["student_id"],pid)).fetchone()
            if not ok: conn.close(); abort(403)
        sql += " AND p.id=?"; params.append(pid)
    if from_date:
        sql += " AND da.session_date>=?"; params.append(from_date)
    if to_date:
        sql += " AND da.session_date<=?"; params.append(to_date)
    sql += " ORDER BY da.session_date,p.sport,s.name"
    rows = conn.execute(sql, params).fetchall(); conn.close()
    out = io.StringIO(); writer = csv.writer(out)
    writer.writerow(["Date","Student ID","Student","Sport","Programme","Attendance Status","Attendance Score","Overall Performance","Coach Remarks","Coach"])
    for r in rows:
        writer.writerow([r["session_date"],r["student_id"],r["name"],r["sport"],r["programme"],r["attendance_status"],r["attendance_score"],r["overall_performance"],r["remarks"],r["coach"]])
    return Response(out.getvalue(), headers=[("Content-Disposition","attachment; filename=inspiria_coaching_report.csv")], content_type="text/csv; charset=utf-8")


@app.route("/reports/projection.csv")
@login_required
def export_projection():
    period = request.args.get("period", "weekly").strip().lower()
    conn = get_db(); projection = build_projection(conn, period); conn.close()
    out = io.StringIO(); writer = csv.writer(out)
    writer.writerow(["Projection Period","From","To","Sport","Athletes","Recorded Sessions","Attendance %","Average Performance","Metric","Metric Average","Metric Samples"])
    for row in projection["sports"]:
        if row["metrics"]:
            for metric in row["metrics"]:
                writer.writerow([projection["period"].title(),projection["from_date"],projection["to_date"],row["sport"],row["athletes"],row["sessions"],row["attendance_rate"],row["avg_performance"],metric["metric_label"],metric["avg_score"],metric["samples"]])
        else:
            writer.writerow([projection["period"].title(),projection["from_date"],projection["to_date"],row["sport"],row["athletes"],row["sessions"],row["attendance_rate"],row["avg_performance"],"","",""])
    filename=f"inspiria_{projection['period']}_sports_projection.csv"
    return Response(out.getvalue(), headers=[("Content-Disposition",f"attachment; filename={filename}")], content_type="text/csv; charset=utf-8")


@app.route("/tournaments")
@app.route("/events")
@login_required
def tournaments():
    exact_date = request.args.get("date", "").strip()
    month = request.args.get("month", "").strip()
    year = request.args.get("year", "").strip()
    conn = get_db(); sql = "SELECT * FROM tournaments WHERE 1=1"; params = []
    if exact_date:
        sql += " AND start_date=?"; params.append(exact_date)
    if year.isdigit():
        sql += " AND substr(start_date,1,4)=?"; params.append(year.zfill(4))
    if month.isdigit():
        sql += " AND substr(start_date,6,2)=?"; params.append(month.zfill(2))
    sql += " ORDER BY start_date"
    ts = conn.execute(sql, params).fetchall()
    today = datetime.now().strftime("%Y-%m-%d")
    upcoming = conn.execute("SELECT * FROM tournaments WHERE start_date>=? ORDER BY start_date LIMIT 6", (today,)).fetchall()
    years = [r[0] for r in conn.execute("SELECT DISTINCT substr(start_date,1,4) y FROM tournaments WHERE start_date IS NOT NULL ORDER BY y DESC").fetchall()]
    conn.close()
    return render_template("tournaments.html", tournaments=ts, upcoming=upcoming, exact_date=exact_date, month=month, year=year, years=years, today=today)


@app.route("/tournament/<int:tournament_id>")
@login_required
def tournament_detail(tournament_id):
    user = current_user(); conn = get_db()
    t = conn.execute("SELECT * FROM tournaments WHERE id=?", (tournament_id,)).fetchone()
    if not t: conn.close(); abort(404)
    department = request.args.get("department", "").strip()
    category = request.args.get("category", "").strip().title()
    match_sql = "SELECT * FROM matches WHERE tournament_id=?"; match_params = [tournament_id]
    if department:
        match_sql += " AND (team_a=? OR team_b=?)"; match_params.extend([department, department])
    if category in ("Boys","Girls"):
        match_sql += " AND category=?"; match_params.append(category)
    match_sql += " ORDER BY match_date,id"
    raw_matches = conn.execute(match_sql, match_params).fetchall()
    matches=[]
    for row in raw_matches:
        item=dict(row)
        item["players_a"]=conn.execute("SELECT * FROM match_players WHERE match_id=? AND side='A' ORDER BY student_name", (row["id"],)).fetchall()
        item["players_b"]=conn.execute("SELECT * FROM match_players WHERE match_id=? AND side='B' ORDER BY student_name", (row["id"],)).fetchall()
        matches.append(item)
    regs = []
    if user["role"] == "admin":
        regs = conn.execute("""SELECT r.*,s.student_id,s.passport_id FROM registrations r JOIN students s ON s.id=r.student_id
                             WHERE r.tournament_id=? ORDER BY r.name,s.name""", (tournament_id,)).fetchall()
    own_registration = None; student = None
    if user["role"] == "student":
        student = conn.execute("SELECT * FROM students WHERE id=?", (user["student_id"],)).fetchone()
        own_registration = conn.execute("SELECT * FROM registrations WHERE tournament_id=? AND student_id=?", (tournament_id,user["student_id"])).fetchone()
    students = conn.execute("SELECT id,name,student_id,passport_id,department FROM students ORDER BY name").fetchall() if user["role"] == "admin" else []
    media = conn.execute("SELECT * FROM event_media WHERE tournament_id=? ORDER BY created_at DESC,id DESC LIMIT 60", (tournament_id,)).fetchall()
    photo_carousel = conn.execute("SELECT * FROM event_media WHERE tournament_id=? AND media_type='photo' ORDER BY created_at DESC,id DESC LIMIT 30", (tournament_id,)).fetchall()
    participants = conn.execute("SELECT * FROM event_participants WHERE tournament_id=? ORDER BY category,department,name", (tournament_id,)).fetchall()
    results = conn.execute("SELECT * FROM tournament_results WHERE tournament_id=? ORDER BY CASE category WHEN 'Boys' THEN 0 ELSE 1 END", (tournament_id,)).fetchall()
    draws = conn.execute("SELECT * FROM draw_documents WHERE tournament_id=? ORDER BY category,format_type,created_at DESC", (tournament_id,)).fetchall()
    templates = conn.execute("SELECT * FROM certificate_templates WHERE tournament_id=? ORDER BY certificate_type", (tournament_id,)).fetchall()
    generated = conn.execute("SELECT * FROM generated_certificates WHERE tournament_id=? ORDER BY certificate_type,recipient_name", (tournament_id,)).fetchall()
    registration_count = conn.execute("SELECT COUNT(*) FROM registrations WHERE tournament_id=?", (tournament_id,)).fetchone()[0]
    conn.close()
    return render_template("tournament_detail.html", t=t, matches=matches, regs=regs, students=students, student=student, own_registration=own_registration,
                           media=media, photo_carousel=photo_carousel, participants=participants, results=results, draws=draws, certificate_templates=templates, generated=generated,
                           department=department, category=category, registration_count=registration_count)

@app.route("/tournament/<int:tournament_id>/register", methods=["POST"])
@admin_required
def register_tournament(tournament_id):
    sid = request.form.get("student_id"); team = request.form.get("team_name", "").strip(); conn = get_db()
    student = conn.execute("SELECT * FROM students WHERE id=?", (sid,)).fetchone()
    if not student:
        conn.close(); abort(404)
    try:
        conn.execute("""INSERT INTO registrations(tournament_id,student_id,team_name,name,department,semester,phone,emergency_contact,blood_group,medical_history)
                      VALUES (?,?,?,?,?,?,?,?,?,?)""",
                     (tournament_id,sid,team,student["name"],student["department"],student["semester"],student["phone"],student["emergency_contact"],student["blood_group"],student["medical_history"]))
        conn.commit(); flash("Registration added.", "success")
    except sqlite3.IntegrityError:
        flash("This student is already registered.", "error")
    finally: conn.close()
    return redirect(url_for("tournament_detail", tournament_id=tournament_id))


@app.route("/tournament/<int:tournament_id>/self-register", methods=["POST"])
@student_required
def self_register_tournament(tournament_id):
    user = current_user(); conn = get_db()
    event = conn.execute("SELECT * FROM tournaments WHERE id=?", (tournament_id,)).fetchone()
    student = conn.execute("SELECT * FROM students WHERE id=?", (user["student_id"],)).fetchone()
    if not event or not student:
        conn.close(); abort(404)
    if not event["registration_open"]:
        conn.close(); flash("Registration is currently closed for this event.", "error"); return redirect(url_for("tournament_detail", tournament_id=tournament_id))
    vals = {k: request.form.get(k, "").strip() for k in ["name","department","semester","phone","emergency_contact","blood_group","medical_history","team_name"]}
    try:
        conn.execute("""UPDATE students SET name=?,department=?,semester=?,phone=?,emergency_contact=?,blood_group=?,medical_history=? WHERE id=?""",
                     (vals["name"],vals["department"],vals["semester"],vals["phone"],vals["emergency_contact"],vals["blood_group"],vals["medical_history"],student["id"]))
        conn.execute("""INSERT INTO registrations(tournament_id,student_id,team_name,name,department,semester,phone,emergency_contact,blood_group,medical_history)
                      VALUES (?,?,?,?,?,?,?,?,?,?)""",
                     (tournament_id,student["id"],vals["team_name"],vals["name"],vals["department"],vals["semester"],vals["phone"],vals["emergency_contact"],vals["blood_group"],vals["medical_history"]))
        conn.commit(); flash("You are registered for this event.", "success")
    except sqlite3.IntegrityError:
        conn.rollback(); flash("You are already registered for this event.", "error")
    finally:
        conn.close()
    return redirect(url_for("tournament_detail", tournament_id=tournament_id))


@app.route("/student/<int:student_id>/performance.csv")
@login_required
def student_performance_export(student_id):
    user = current_user(); conn = get_db()
    if not can_access_student(conn, student_id, user):
        conn.close(); abort(403)
    rows = conn.execute("""SELECT da.session_date,p.sport,p.name programme,da.attendance_status,da.attendance_score,da.overall_performance,da.remarks,c.name coach
                         FROM daily_assessments da JOIN enrollments e ON e.id=da.enrollment_id JOIN programmes p ON p.id=e.programme_id
                         LEFT JOIN coaches c ON c.id=da.coach_id WHERE e.student_id=? ORDER BY da.session_date DESC""", (student_id,)).fetchall()
    student = conn.execute("SELECT student_id,name FROM students WHERE id=?", (student_id,)).fetchone(); conn.close()
    out=io.StringIO(); w=csv.writer(out); w.writerow(["Date","Sport","Programme","Attendance","Attendance Rating","Overall Performance","Coach Remarks","Coach"])
    for r in rows: w.writerow([r["session_date"],r["sport"],r["programme"],r["attendance_status"],r["attendance_score"],r["overall_performance"],r["remarks"],r["coach"]])
    filename=f"{student['student_id']}_performance.csv" if student else "performance.csv"
    return Response(out.getvalue(), headers=[("Content-Disposition",f"attachment; filename={filename}")], content_type="text/csv; charset=utf-8")


@app.route("/matches")
@login_required
def matches():
    conn = get_db(); raw = conn.execute("""SELECT m.*,t.name tournament,t.sport FROM matches m JOIN tournaments t ON t.id=m.tournament_id
                         ORDER BY CASE m.status WHEN 'Live' THEN 0 WHEN 'Scheduled' THEN 1 ELSE 2 END,m.match_date""").fetchall()
    rows=[]
    for r in raw:
        item=dict(r); item["stats"]=conn.execute("SELECT stat_key,stat_label,stat_value FROM match_live_stats WHERE match_id=? ORDER BY id", (r["id"],)).fetchall(); item["stat_map"]={x["stat_key"]:x["stat_value"] for x in item["stats"]}; rows.append(item)
    conn.close(); return render_template("matches.html", matches=rows)


@app.route("/match/<int:match_id>/score", methods=["POST"])
@admin_required
def update_score(match_id):
    a=request.form.get("score_a","0"); b=request.form.get("score_b","0"); status=request.form.get("status","Scheduled"); note=request.form.get("live_note","").strip()
    conn=get_db(); match=conn.execute("""SELECT m.*,t.sport FROM matches m JOIN tournaments t ON t.id=m.tournament_id WHERE m.id=?""", (match_id,)).fetchone()
    if not match: conn.close(); abort(404)
    photo_path=match["photo_path"]
    try:
        uploaded=request.files.get("photo")
        if uploaded and uploaded.filename: photo_path=save_upload(uploaded,"matches",ALLOWED_IMAGE_EXTENSIONS)
    except ValueError as e:
        flash(str(e),"error")
    conn.execute("UPDATE matches SET score_a=?,score_b=?,status=?,live_note=?,photo_path=? WHERE id=?",(a,b,status,note,photo_path,match_id))
    for key,label in LIVE_SCORE_FIELDS.get(match["sport"], []):
        value=request.form.get(f"stat_{key}","").strip()
        conn.execute("""INSERT INTO match_live_stats(match_id,stat_key,stat_label,stat_value) VALUES (?,?,?,?)
                      ON CONFLICT(match_id,stat_key) DO UPDATE SET stat_label=excluded.stat_label,stat_value=excluded.stat_value""", (match_id,key,label,value))
    conn.commit(); conn.close(); return redirect(request.referrer or url_for("matches"))


@app.route("/api/matches")
def api_matches():
    conn=get_db(); rows=conn.execute("""SELECT m.id,t.name tournament,t.sport,m.round_name,m.match_date,m.venue,m.team_a,m.team_b,m.score_a,m.score_b,m.status,m.live_note,m.photo_path
                         FROM matches m JOIN tournaments t ON t.id=m.tournament_id WHERE m.status IN ('Live','Scheduled') ORDER BY m.match_date""").fetchall()
    result=[]
    for r in rows:
        item=dict(r); item["stats"]=[dict(x) for x in conn.execute("SELECT stat_key,stat_label,stat_value FROM match_live_stats WHERE match_id=?", (r["id"],)).fetchall()]; result.append(item)
    conn.close(); return jsonify(result)


@app.route("/chat", methods=["GET", "POST"])
@login_required
def chat():
    user=current_user(); conn=get_db()
    if user["role"]=="admin":
        contacts=conn.execute("SELECT id,display_name,designation,role FROM users WHERE active=1 AND id<>? ORDER BY role,display_name", (user["id"],)).fetchall()
    elif user["role"]=="coach":
        contacts=conn.execute("""SELECT DISTINCT u.id,u.display_name,u.designation,u.role FROM users u JOIN students s ON s.id=u.student_id
                    JOIN enrollments e ON e.student_id=s.id JOIN programmes p ON p.id=e.programme_id
                    WHERE u.active=1 AND u.role='student' AND p.coach_id=? AND e.status='Active' ORDER BY u.display_name""", (user["coach_id"],)).fetchall()
    else:
        contacts=conn.execute("SELECT id,display_name,designation,role FROM users WHERE active=1 AND role='coach' ORDER BY display_name").fetchall()
    allowed={c["id"] for c in contacts}
    selected_id=request.args.get("with", type=int)
    if request.method=="POST":
        selected_id=request.form.get("receiver_id", type=int); message=request.form.get("message","").strip()
        if selected_id not in allowed:
            conn.close(); abort(403)
        if message:
            conn.execute("INSERT INTO messages(sender_user_id,receiver_user_id,message) VALUES (?,?,?)", (user["id"],selected_id,message)); conn.commit()
        conn.close(); return redirect(url_for("chat", **{"with":selected_id}))
    conversation=[]; selected=None
    if selected_id in allowed:
        selected=conn.execute("SELECT id,display_name,designation,role FROM users WHERE id=?", (selected_id,)).fetchone()
        conversation=conn.execute("""SELECT m.*,su.display_name sender_name FROM messages m JOIN users su ON su.id=m.sender_user_id
                            WHERE (m.sender_user_id=? AND m.receiver_user_id=?) OR (m.sender_user_id=? AND m.receiver_user_id=?)
                            ORDER BY m.created_at,m.id""", (user["id"],selected_id,selected_id,user["id"])).fetchall()
        conn.execute("UPDATE messages SET read_at=CURRENT_TIMESTAMP WHERE receiver_user_id=? AND sender_user_id=? AND read_at IS NULL", (user["id"],selected_id)); conn.commit()
    conn.close(); return render_template("chat.html", contacts=contacts, selected=selected, conversation=conversation)


@app.route("/gallery")
@login_required
def gallery():
    event_id=request.args.get("event", type=int); media_type=request.args.get("type","").strip()
    conn=get_db(); sql="""SELECT em.*,t.name event_name,t.sport FROM event_media em LEFT JOIN tournaments t ON t.id=em.tournament_id WHERE 1=1"""; params=[]
    if event_id: sql+=" AND em.tournament_id=?"; params.append(event_id)
    if media_type in ("photo","video"): sql+=" AND em.media_type=?"; params.append(media_type)
    sql+=" ORDER BY em.created_at DESC,em.id DESC"
    media=conn.execute(sql,params).fetchall(); events=conn.execute("SELECT id,name,start_date FROM tournaments ORDER BY start_date DESC").fetchall(); conn.close()
    return render_template("gallery.html", media=media, events=events, event_id=event_id, media_type=media_type)


@app.route("/gallery/<int:media_id>/download")
@login_required
def download_media(media_id):
    conn=get_db(); row=conn.execute("SELECT * FROM event_media WHERE id=?", (media_id,)).fetchone(); conn.close()
    if not row: abort(404)
    full=os.path.join(app.static_folder,row["file_path"]); directory=os.path.dirname(full); filename=os.path.basename(full)
    if not os.path.exists(full): abort(404)
    return send_from_directory(directory, filename, as_attachment=True)


@app.route("/admin/media/upload", methods=["POST"])
@admin_required
def upload_event_media():
    event_id=request.form.get("tournament_id", type=int); caption=request.form.get("caption","").strip(); upload=request.files.get("file")
    if not upload or not upload.filename:
        flash("Choose a photo or video to upload.","error"); return redirect(request.referrer or url_for("gallery"))
    ext=upload.filename.rsplit(".",1)[-1].lower() if "." in upload.filename else ""; media_type="photo" if ext in ALLOWED_IMAGE_EXTENSIONS else "video" if ext in ALLOWED_VIDEO_EXTENSIONS else None
    if not media_type:
        flash("Supported media: JPG, PNG, WEBP, MP4, WEBM or MOV.","error"); return redirect(request.referrer or url_for("gallery"))
    try:
        path=save_upload(upload,"events",ALLOWED_IMAGE_EXTENSIONS|ALLOWED_VIDEO_EXTENSIONS)
        user=current_user(); conn=get_db(); conn.execute("INSERT INTO event_media(tournament_id,media_type,file_path,caption,uploaded_by) VALUES (?,?,?,?,?)", (event_id,media_type,path,caption,user["id"])); conn.commit(); conn.close(); flash("Event media uploaded.","success")
    except ValueError as e: flash(str(e),"error")
    return redirect(request.referrer or url_for("gallery"))


@app.route("/admin/tournament/<int:tournament_id>/poster", methods=["POST"])
@admin_required
def upload_event_poster(tournament_id):
    upload=request.files.get("poster")
    try:
        path=save_upload(upload,"posters",ALLOWED_IMAGE_EXTENSIONS)
        if not path: raise ValueError("Choose a poster image.")
        conn=get_db(); conn.execute("UPDATE tournaments SET poster_path=? WHERE id=?", (path,tournament_id)); conn.commit(); conn.close(); flash("Event poster updated.","success")
    except ValueError as e: flash(str(e),"error")
    return redirect(request.referrer or url_for("tournament_detail", tournament_id=tournament_id))


@app.route("/admin/fixture", methods=["POST"])
@admin_required
def add_fixture():
    vals=[request.form.get(k,"").strip() for k in ["tournament_id","round_name","match_date","venue","team_a","team_b"]]; category=request.form.get("category","Boys").strip().title()
    if not vals[0] or not vals[4] or not vals[5]:
        flash("Event and both departments/teams are required.","error"); return redirect(request.referrer or url_for("admin"))
    conn=get_db(); conn.execute("INSERT INTO matches(tournament_id,round_name,match_date,venue,team_a,team_b,category) VALUES (?,?,?,?,?,?,?)", (*vals,category)); conn.commit(); conn.close(); flash("Fixture created.","success"); return redirect(request.referrer or url_for("admin"))


@app.route("/admin/match/<int:match_id>/delete", methods=["POST"])
@admin_required
def delete_fixture(match_id):
    conn=get_db(); row=conn.execute("SELECT tournament_id FROM matches WHERE id=?", (match_id,)).fetchone()
    if row: conn.execute("DELETE FROM matches WHERE id=?", (match_id,)); conn.commit()
    conn.close(); flash("Fixture deleted.","success")
    return redirect(request.referrer or (url_for("tournament_detail", tournament_id=row["tournament_id"]) if row else url_for("admin")))


@app.route("/admin/match/<int:match_id>/player", methods=["POST"])
@admin_required
def add_match_player(match_id):
    side=request.form.get("side","A").strip().upper(); student_id=request.form.get("student_id","").strip(); typed_name=request.form.get("student_name","").strip()
    if side not in ("A","B"):
        flash("Choose Team A or Team B.","error"); return redirect(request.referrer or url_for("matches"))
    conn=get_db(); match=conn.execute("SELECT * FROM matches WHERE id=?",(match_id,)).fetchone()
    if not match: conn.close(); abort(404)
    count=conn.execute("SELECT COUNT(*) FROM match_players WHERE match_id=? AND side=?",(match_id,side)).fetchone()[0]
    if count >= 15:
        conn.close(); flash("A fixture squad can contain a maximum of 15 students per team.","error"); return redirect(request.referrer or url_for("matches"))
    expected_department=match["team_a"] if side=="A" else match["team_b"]
    student=None
    if student_id.isdigit():
        student=conn.execute("SELECT id,name,department FROM students WHERE id=?",(int(student_id),)).fetchone()
        if student and not department_match(student["department"], expected_department):
            conn.close(); flash(f"Select a student from {expected_department} for this side.","error"); return redirect(request.referrer or url_for("matches"))
    if student:
        name=student["name"]; department=student["department"]; sid=student["id"]
    else:
        name=typed_name; department=expected_department; sid=None
    if not name:
        conn.close(); flash("Choose a student or enter a player name.","error"); return redirect(request.referrer or url_for("matches"))
    duplicate=conn.execute("SELECT 1 FROM match_players WHERE match_id=? AND side=? AND lower(student_name)=lower(?)",(match_id,side,name)).fetchone()
    if duplicate:
        conn.close(); flash("That student is already in this team list.","error"); return redirect(request.referrer or url_for("matches"))
    conn.execute("INSERT INTO match_players(match_id,side,student_name,department,student_id) VALUES (?,?,?,?,?)",(match_id,side,name,department,sid)); conn.commit(); conn.close()
    flash(f"Player added to {expected_department} squad ({count+1}/15).","success"); return redirect(request.referrer or url_for("matches"))


@app.route("/admin/match/player/<int:player_id>/delete", methods=["POST"])
@admin_required
def delete_match_player(player_id):
    conn=get_db(); conn.execute("DELETE FROM match_players WHERE id=?",(player_id,)); conn.commit(); conn.close(); flash("Player removed from fixture roster.","success")
    return redirect(request.referrer or url_for("matches"))


@app.route("/admin/tournament/<int:tournament_id>/participants/import", methods=["POST"])
@admin_required
def import_event_participants(tournament_id):
    upload=request.files.get("file")
    try:
        rows=participant_sheet_rows(upload)
        conn=get_db(); added=0; skipped=0
        for row in rows:
            student=conn.execute("SELECT id FROM students WHERE lower(name)=lower(?) AND department=? LIMIT 1", (row["name"],row["department"])).fetchone()
            try:
                conn.execute("INSERT INTO event_participants(tournament_id,student_id,name,department,category,year) VALUES (?,?,?,?,?,?)",
                             (tournament_id,student["id"] if student else None,row["name"],row["department"],row["category"],row["year"])); added+=1
            except sqlite3.IntegrityError: skipped+=1
        conn.commit(); conn.close(); flash(f"Official participant list updated: {added} added, {skipped} duplicates skipped.","success")
    except ValueError as e: flash(str(e),"error")
    return redirect(url_for("tournament_detail",tournament_id=tournament_id))


@app.route("/admin/participant/<int:participant_id>/delete", methods=["POST"])
@admin_required
def delete_event_participant(participant_id):
    conn=get_db(); row=conn.execute("SELECT tournament_id FROM event_participants WHERE id=?",(participant_id,)).fetchone()
    if row: conn.execute("DELETE FROM event_participants WHERE id=?",(participant_id,)); conn.commit()
    conn.close(); flash("Participant removed from official list.","success")
    return redirect(request.referrer or (url_for("tournament_detail",tournament_id=row["tournament_id"]) if row else url_for("tournaments")))


@app.route("/admin/tournament/<int:tournament_id>/result", methods=["POST"])
@admin_required
def save_tournament_result(tournament_id):
    category=request.form.get("category","Boys").strip().title(); winner=request.form.get("winner","").strip(); runner=request.form.get("runner_up","").strip(); notes=request.form.get("notes","").strip()
    if category not in ("Boys","Girls") or not winner or not runner:
        flash("Category, Winner and Runner-up are required.","error"); return redirect(request.referrer or url_for("tournament_detail",tournament_id=tournament_id))
    conn=get_db(); conn.execute("""INSERT INTO tournament_results(tournament_id,category,winner,runner_up,notes) VALUES (?,?,?,?,?)
                                ON CONFLICT(tournament_id,category) DO UPDATE SET winner=excluded.winner,runner_up=excluded.runner_up,notes=excluded.notes""",
                                (tournament_id,category,winner,runner,notes)); conn.commit(); conn.close(); flash(f"{category} result saved.","success")
    return redirect(url_for("tournament_detail",tournament_id=tournament_id))


@app.route("/admin/result/<int:result_id>/delete", methods=["POST"])
@admin_required
def delete_tournament_result(result_id):
    conn=get_db(); conn.execute("DELETE FROM tournament_results WHERE id=?",(result_id,)); conn.commit(); conn.close(); flash("Result removed.","success")
    return redirect(request.referrer or url_for("tournaments"))


@app.route("/admin/tournament/<int:tournament_id>/draw", methods=["POST"])
@admin_required
def upload_draw_document(tournament_id):
    category=request.form.get("category","Boys").strip().title(); format_type=request.form.get("format_type","Knockout").strip().title(); upload=request.files.get("file")
    if category not in ("Boys","Girls") or format_type not in ("League","Knockout"):
        flash("Choose Boys/Girls and League/Knockout.","error"); return redirect(request.referrer or url_for("tournament_detail",tournament_id=tournament_id))
    try:
        path=save_upload(upload,"draws",ALLOWED_DRAW_EXTENSIONS)
        if not path: raise ValueError("Choose a PDF draw file.")
        user=current_user(); conn=get_db(); conn.execute("INSERT INTO draw_documents(tournament_id,category,format_type,file_path,uploaded_by) VALUES (?,?,?,?,?)",(tournament_id,category,format_type,path,user["id"])); conn.commit(); conn.close(); flash("Fixture projection PDF uploaded.","success")
    except ValueError as e: flash(str(e),"error")
    return redirect(url_for("tournament_detail",tournament_id=tournament_id))


@app.route("/admin/draw/<int:draw_id>/delete", methods=["POST"])
@admin_required
def delete_draw_document(draw_id):
    conn=get_db(); row=conn.execute("SELECT * FROM draw_documents WHERE id=?",(draw_id,)).fetchone()
    if row:
        full=os.path.join(app.static_folder,row["file_path"])
        if os.path.exists(full): os.remove(full)
        conn.execute("DELETE FROM draw_documents WHERE id=?",(draw_id,)); conn.commit()
    conn.close(); flash("Draw document deleted.","success"); return redirect(request.referrer or url_for("tournaments"))


@app.route("/announcements")
@login_required
def announcements():
    conn=get_db(); rows=conn.execute("""SELECT a.*,u.display_name author FROM announcements a LEFT JOIN users u ON u.id=a.created_by
                                      WHERE a.active=1 OR ?='admin' ORDER BY a.created_at DESC,a.id DESC""",(current_user()["role"],)).fetchall(); conn.close()
    return render_template("announcements.html", announcements=rows)


@app.route("/admin/announcement", methods=["POST"])
@admin_required
def add_announcement():
    title=request.form.get("title","").strip(); message=request.form.get("message","").strip(); upload=request.files.get("image")
    if not title or not message:
        flash("Announcement title and message are required.","error"); return redirect(request.referrer or url_for("announcements"))
    image_path=None
    try:
        if upload and upload.filename:
            image_path=save_upload(upload,"announcements",ALLOWED_IMAGE_EXTENSIONS)
    except ValueError as exc:
        flash(str(exc),"error"); return redirect(request.referrer or url_for("announcements"))
    user=current_user(); conn=get_db(); conn.execute("INSERT INTO announcements(title,message,image_path,created_by) VALUES (?,?,?,?)",(title,message,image_path,user["id"])); conn.commit(); conn.close(); flash("Announcement shared with everyone.","success")
    return redirect(url_for("announcements"))


@app.route("/admin/announcement/<int:announcement_id>/delete", methods=["POST"])
@admin_required
def delete_announcement(announcement_id):
    conn=get_db(); row=conn.execute("SELECT image_path FROM announcements WHERE id=?",(announcement_id,)).fetchone()
    if row:
        conn.execute("DELETE FROM announcements WHERE id=?",(announcement_id,)); conn.commit(); remove_static_file(row["image_path"])
    conn.close(); flash("Announcement deleted.","success"); return redirect(url_for("announcements"))


@app.route("/admin/media/<int:media_id>/delete", methods=["POST"])
@admin_required
def delete_event_media(media_id):
    conn=get_db(); row=conn.execute("SELECT * FROM event_media WHERE id=?",(media_id,)).fetchone()
    if row:
        full=os.path.join(app.static_folder,row["file_path"])
        if os.path.exists(full): os.remove(full)
        conn.execute("DELETE FROM event_media WHERE id=?",(media_id,)); conn.commit()
    conn.close(); flash("Event media deleted.","success"); return redirect(request.referrer or url_for("gallery"))


@app.route("/admin/tournament/<int:tournament_id>/certificate-template", methods=["POST"])
@admin_required
def upload_certificate_template(tournament_id):
    certificate_type=request.form.get("certificate_type","Participation").strip().title(); upload=request.files.get("file")
    if certificate_type not in ("Participation","Achievement"):
        flash("Choose Participation or Achievement template.","error"); return redirect(request.referrer or url_for("tournament_detail",tournament_id=tournament_id))
    try:
        path=save_upload(upload,"certificates/templates",ALLOWED_CERTIFICATE_EXTENSIONS)
        if not path: raise ValueError("Choose a PDF, JPG or PNG certificate template.")
        user=current_user(); conn=get_db(); old=conn.execute("SELECT file_path FROM certificate_templates WHERE tournament_id=? AND certificate_type=?",(tournament_id,certificate_type)).fetchone()
        conn.execute("""INSERT INTO certificate_templates(tournament_id,certificate_type,file_path,uploaded_by) VALUES (?,?,?,?)
                      ON CONFLICT(tournament_id,certificate_type) DO UPDATE SET file_path=excluded.file_path,uploaded_by=excluded.uploaded_by,created_at=CURRENT_TIMESTAMP""",
                     (tournament_id,certificate_type,path,user["id"])); conn.commit(); conn.close()
        if old:
            old_full=os.path.join(app.static_folder,old["file_path"])
            if os.path.exists(old_full): os.remove(old_full)
        flash(f"{certificate_type} certificate template uploaded.","success")
    except ValueError as e: flash(str(e),"error")
    return redirect(url_for("tournament_detail",tournament_id=tournament_id))


@app.route("/admin/tournament/<int:tournament_id>/certificates/generate", methods=["POST"])
@admin_required
def generate_event_certificates(tournament_id):
    certificate_type=request.form.get("certificate_type","Participation").strip().title()
    conn=get_db(); event=conn.execute("SELECT * FROM tournaments WHERE id=?",(tournament_id,)).fetchone(); template=conn.execute("SELECT * FROM certificate_templates WHERE tournament_id=? AND certificate_type=?",(tournament_id,certificate_type)).fetchone()
    if not event or not template:
        conn.close(); flash(f"Upload the {certificate_type} certificate template first.","error"); return redirect(url_for("tournament_detail",tournament_id=tournament_id))
    participants=conn.execute("SELECT * FROM event_participants WHERE tournament_id=? ORDER BY id",(tournament_id,)).fetchall(); selected=[]
    if certificate_type=="Participation":
        selected=[(p,None) for p in participants]
    else:
        results=conn.execute("SELECT * FROM tournament_results WHERE tournament_id=?",(tournament_id,)).fetchall()
        for p in participants:
            for r in results:
                if p["category"]!=r["category"]: continue
                p_name=p["name"].strip().lower(); p_dept=p["department"].strip().lower()
                if p_name==r["winner"].strip().lower() or p_dept==r["winner"].strip().lower(): selected.append((p,"Winner")); break
                if p_name==r["runner_up"].strip().lower() or p_dept==r["runner_up"].strip().lower(): selected.append((p,"Runner-up")); break
    if not selected:
        conn.close(); flash("No eligible official participants were found for certificate generation.","error"); return redirect(url_for("tournament_detail",tournament_id=tournament_id))
    template_full=os.path.join(app.static_folder,template["file_path"]); generated_count=0
    for serial,(p,achievement) in enumerate(selected,1):
        existing=conn.execute("SELECT id FROM generated_certificates WHERE tournament_id=? AND participant_id=? AND certificate_type=?",(tournament_id,p["id"],certificate_type)).fetchone()
        if existing: continue
        cert_no=next_certificate_no(tournament_id,certificate_type,conn.execute("SELECT COUNT(*)+1 FROM generated_certificates WHERE tournament_id=? AND certificate_type=?",(tournament_id,certificate_type)).fetchone()[0])
        filename=f"{uuid.uuid4().hex}.pdf"; rel=f"uploads/certificates/generated/{filename}"; full=os.path.join(app.static_folder,rel)
        try:
            create_certificate_pdf(template_full,full,p["name"],event["name"],p["category"],p["year"],certificate_type,cert_no,achievement)
        except Exception as exc:
            conn.rollback(); conn.close(); flash(f"Certificate generation failed: {exc}","error"); return redirect(url_for("tournament_detail",tournament_id=tournament_id))
        conn.execute("""INSERT INTO generated_certificates(tournament_id,participant_id,student_id,certificate_type,recipient_name,achievement,certificate_no,file_path)
                      VALUES (?,?,?,?,?,?,?,?)""",(tournament_id,p["id"],p["student_id"],certificate_type,p["name"],achievement,cert_no,rel)); generated_count+=1
    conn.commit(); conn.close(); flash(f"Generated {generated_count} new {certificate_type.lower()} certificates.","success")
    return redirect(url_for("tournament_detail",tournament_id=tournament_id))


@app.route("/certificate/<int:certificate_id>/download")
@login_required
def download_generated_certificate(certificate_id):
    user=current_user(); conn=get_db(); row=conn.execute("SELECT * FROM generated_certificates WHERE id=?",(certificate_id,)).fetchone(); conn.close()
    if not row: abort(404)
    if user["role"]=="student" and row["student_id"]!=user["student_id"]: abort(403)
    if user["role"]=="coach": abort(403)
    full=os.path.join(app.static_folder,row["file_path"])
    if not os.path.exists(full): abort(404)
    return send_from_directory(os.path.dirname(full),os.path.basename(full),as_attachment=True,download_name=f"{row['certificate_no']}.pdf")


@app.route("/admin/sport/<int:sport_id>/cover", methods=["POST"])
@admin_required
def upload_sport_cover(sport_id):
    upload=request.files.get("photo")
    try:
        path=save_upload(upload,"facilities",ALLOWED_IMAGE_EXTENSIONS)
        if not path: raise ValueError("Choose a facility photo.")
        conn=get_db(); old=conn.execute("SELECT cover_path FROM sports WHERE id=?",(sport_id,)).fetchone(); conn.execute("UPDATE sports SET cover_path=? WHERE id=?",(path,sport_id)); conn.commit(); conn.close()
        if old and old["cover_path"]:
            old_full=os.path.join(app.static_folder,old["cover_path"])
            if os.path.exists(old_full): os.remove(old_full)
        flash("Facility cover photo updated.","success")
    except ValueError as e: flash(str(e),"error")
    return redirect(request.referrer or url_for("sport_detail",sport_id=sport_id))


@app.route("/admin/sport/<int:sport_id>/cover/delete", methods=["POST"])
@admin_required
def delete_sport_cover(sport_id):
    conn=get_db(); row=conn.execute("SELECT cover_path FROM sports WHERE id=?",(sport_id,)).fetchone()
    if row and row["cover_path"]:
        full=os.path.join(app.static_folder,row["cover_path"])
        if os.path.exists(full): os.remove(full)
        conn.execute("UPDATE sports SET cover_path=NULL WHERE id=?",(sport_id,)); conn.commit()
    conn.close(); flash("Facility cover photo removed.","success"); return redirect(request.referrer or url_for("sport_detail",sport_id=sport_id))


@app.route("/admin/sport/<int:sport_id>/facility", methods=["POST"])
@admin_required
def update_sport_facility(sport_id):
    venue=request.form.get("venue","").strip()
    conn=get_db(); conn.execute("UPDATE sports SET venue=? WHERE id=?",(venue,sport_id)); conn.commit(); conn.close(); flash("Facility name updated.","success")
    return redirect(url_for("sport_detail",sport_id=sport_id))


@app.route("/admin/sport/<int:sport_id>/media", methods=["POST"])
@admin_required
def upload_facility_media(sport_id):
    upload=request.files.get("photo"); caption=request.form.get("caption","").strip()
    try:
        path=save_upload(upload,"facilities",ALLOWED_IMAGE_EXTENSIONS)
        if not path: raise ValueError("Choose a facility photo.")
        user=current_user(); conn=get_db(); conn.execute("INSERT INTO facility_media(sport_id,file_path,caption,uploaded_by) VALUES (?,?,?,?)",(sport_id,path,caption,user["id"])); conn.commit(); conn.close(); flash("Facility photo added.","success")
    except ValueError as e: flash(str(e),"error")
    return redirect(url_for("sport_detail",sport_id=sport_id))


@app.route("/admin/facility-media/<int:media_id>/delete", methods=["POST"])
@admin_required
def delete_facility_media(media_id):
    conn=get_db(); row=conn.execute("SELECT * FROM facility_media WHERE id=?",(media_id,)).fetchone()
    if row:
        full=os.path.join(app.static_folder,row["file_path"])
        if os.path.exists(full): os.remove(full)
        conn.execute("DELETE FROM facility_media WHERE id=?",(media_id,)); conn.commit()
    conn.close(); flash("Facility photo deleted.","success"); return redirect(request.referrer or url_for("sports"))


@app.route("/admin/coach/<int:coach_id>/update", methods=["POST"])
@admin_required
def admin_update_coach(coach_id):
    name=request.form.get("name","").strip(); sport=request.form.get("sport","").strip(); email=request.form.get("email","").strip(); username=request.form.get("username","").strip().lower()
    conn=get_db(); conn.execute("UPDATE coaches SET name=?,sport=?,email=?,username=? WHERE id=?", (name,sport,email,username,coach_id)); conn.execute("UPDATE users SET display_name=?,username=?,designation=? WHERE coach_id=?", (name,username,f"{sport} Coach",coach_id)); conn.commit(); conn.close(); flash("Coach profile updated.","success"); return redirect(url_for("admin"))


@app.route("/admin/programme/<int:programme_id>/update", methods=["POST"])
@admin_required
def admin_update_programme(programme_id):
    vals=[request.form.get(k,"").strip() for k in ["name","sport","coach_id","schedule","venue"]]; active=1 if request.form.get("active") else 0
    coach_id=int(vals[2]) if vals[2].isdigit() else None
    conn=get_db(); conn.execute("UPDATE programmes SET name=?,sport=?,coach_id=?,schedule=?,venue=?,active=? WHERE id=?", (vals[0],vals[1],coach_id,vals[3],vals[4],active,programme_id)); conn.commit(); conn.close(); flash("Programme updated.","success"); return redirect(url_for("admin"))


@app.route("/admin/tournament/<int:tournament_id>/update", methods=["POST"])
@admin_required
def admin_update_tournament(tournament_id):
    vals=[request.form.get(k,"").strip() for k in ["name","sport","start_date","end_date","venue","status","description"]]; reg=1 if request.form.get("registration_open") else 0
    conn=get_db(); conn.execute("UPDATE tournaments SET name=?,sport=?,start_date=?,end_date=?,venue=?,status=?,description=?,registration_open=? WHERE id=?", (*vals,reg,tournament_id)); conn.commit(); conn.close(); flash("Event updated.","success"); return redirect(request.referrer or url_for("admin"))


@app.route("/admin/match/<int:match_id>/update", methods=["POST"])
@admin_required
def admin_update_match(match_id):
    vals=[request.form.get(k,"").strip() for k in ["round_name","match_date","venue","team_a","team_b","score_a","score_b","status"]]
    category=request.form.get("category","Boys").strip().title()
    conn=get_db(); conn.execute("UPDATE matches SET round_name=?,match_date=?,venue=?,team_a=?,team_b=?,score_a=?,score_b=?,status=?,category=? WHERE id=?", (*vals,category,match_id)); conn.commit(); conn.close(); flash("Fixture and score updated.","success"); return redirect(request.referrer or url_for("admin"))


@app.route("/admin/user/<int:user_id>/toggle", methods=["POST"])
@admin_required
def admin_toggle_user(user_id):
    me=current_user()
    if user_id==me["id"]:
        flash("You cannot disable your own active session.","error"); return redirect(url_for("admin"))
    conn=get_db(); row=conn.execute("SELECT active FROM users WHERE id=?", (user_id,)).fetchone()
    if row: conn.execute("UPDATE users SET active=? WHERE id=?", (0 if row["active"] else 1,user_id)); conn.commit()
    conn.close(); return redirect(url_for("admin"))


@app.route("/admin/user/<int:user_id>/password", methods=["POST"])
@admin_required
def admin_reset_password(user_id):
    password=request.form.get("password","")
    if len(password)<8:
        flash("New password must contain at least 8 characters.","error"); return redirect(url_for("admin"))
    conn=get_db(); conn.execute("UPDATE users SET password_hash=? WHERE id=?", (generate_password_hash(password),user_id)); conn.commit(); conn.close(); flash("Account password reset.","success"); return redirect(url_for("admin"))


@app.route("/admin/sport/<int:sport_id>/update", methods=["POST"])
@admin_required
def admin_update_sport(sport_id):
    name=request.form.get("name","").strip(); activity_type=request.form.get("activity_type","Activity").strip(); venue=request.form.get("venue","").strip(); active=1 if request.form.get("is_active") else 0
    conn=get_db(); conn.execute("UPDATE sports SET name=?,activity_type=?,venue=?,is_active=? WHERE id=?", (name,activity_type,venue,active,sport_id)); conn.commit(); conn.close(); flash("Sports activity updated.","success"); return redirect(url_for("admin"))


@app.route("/admin")
@admin_required
def admin():
    conn=get_db()
    students=conn.execute("SELECT * FROM students ORDER BY id DESC").fetchall()
    tournaments=conn.execute("SELECT * FROM tournaments ORDER BY start_date DESC,id DESC").fetchall()
    coaches=conn.execute("SELECT * FROM coaches ORDER BY sport").fetchall()
    programmes=conn.execute("SELECT p.*,c.name coach FROM programmes p LEFT JOIN coaches c ON c.id=p.coach_id ORDER BY p.sport").fetchall()
    users=conn.execute("SELECT id,username,display_name,designation,role,active FROM users ORDER BY role,display_name").fetchall()
    matches=conn.execute("""SELECT m.*,t.name tournament,t.sport FROM matches m JOIN tournaments t ON t.id=m.tournament_id ORDER BY m.match_date DESC,m.id DESC LIMIT 30""").fetchall()
    sports=conn.execute("SELECT * FROM sports ORDER BY name").fetchall()
    announcements_count=conn.execute("SELECT COUNT(*) FROM announcements WHERE active=1").fetchone()[0]
    participant_count=conn.execute("SELECT COUNT(*) FROM event_participants").fetchone()[0]
    conn.close()
    return render_template("admin.html",students=students,tournaments=tournaments,coaches=coaches,programmes=programmes,users=users,matches=matches,sports=sports,announcements_count=announcements_count,participant_count=participant_count)


@app.route("/admin/tournament", methods=["POST"])
@admin_required
def add_tournament():
    fields=["name","sport","start_date","end_date","venue","status","description"]
    vals=[request.form.get(x,"").strip() for x in fields]
    conn=get_db(); conn.execute("INSERT INTO tournaments(name,sport,start_date,end_date,venue,status,description) VALUES (?,?,?,?,?,?,?)",vals); conn.commit(); conn.close()
    flash("Tournament created.","success"); return redirect(url_for("admin"))


@app.errorhandler(403)
def forbidden(_):
    return render_template("403.html"), 403


init_db()

if __name__ == "__main__":
    app.run(host="0.0.0.0", port=5000, debug=True)
