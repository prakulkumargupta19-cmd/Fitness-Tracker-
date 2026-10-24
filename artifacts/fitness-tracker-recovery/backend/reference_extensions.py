"""Complete the Fit-Track operations that are not part of root authentication.

This module deliberately sits beside ``reference_features``.  It adds the
reference project's bulk operations, exports, import workflows, and binary
file endpoints without replacing the existing login or OTP implementation.
"""

from datetime import date, datetime, timedelta
import io
import os
from pathlib import Path
from uuid import uuid4

from flask import Response, jsonify, request, send_file, session
from werkzeug.utils import secure_filename

try:
    from backend.db import get_db
    from backend.auth.decorators import role_required
    from backend.reference_features import (
        _coach_can_access_enrollment,
        _coach_can_access_programme,
        _coach_can_access_student,
        _commit,
        _row,
        _rows,
    )
except ModuleNotFoundError:
    from db import get_db
    from auth.decorators import role_required
    from reference_features import (
        _coach_can_access_enrollment,
        _coach_can_access_programme,
        _coach_can_access_student,
        _commit,
        _row,
        _rows,
    )


ADMIN = role_required("admin")
STAFF = role_required("admin", "coach")
ALL_USERS = role_required("admin", "coach", "student")

UPLOAD_ROOT = Path(
    os.getenv(
        "FIT_TRACK_UPLOAD_ROOT",
        Path(__file__).resolve().parent.parent / "static" / "uploads",
    )
)
IMAGE_EXTENSIONS = {"png", "jpg", "jpeg", "webp"}
VIDEO_EXTENSIONS = {"mp4", "webm", "mov"}
DOCUMENT_EXTENSIONS = {"pdf"}
CERTIFICATE_EXTENSIONS = {"pdf", "png", "jpg", "jpeg"}


def _required(data, *names):
    missing = [name for name in names if data.get(name) in (None, "")]
    if missing:
        return jsonify(success=False, message=f"Required fields: {', '.join(missing)}."), 400
    return None


def _body():
    return request.get_json(silent=True) or {}


def _upload(file_storage, folder, allowed_extensions):
    if not file_storage or not file_storage.filename:
        raise ValueError("A file is required.")
    filename = secure_filename(file_storage.filename)
    if "." not in filename:
        raise ValueError("The uploaded file has no extension.")
    extension = filename.rsplit(".", 1)[1].lower()
    if extension not in allowed_extensions:
        raise ValueError(f"Unsupported file type: .{extension}.")
    target_dir = UPLOAD_ROOT / folder
    target_dir.mkdir(parents=True, exist_ok=True)
    target = target_dir / f"{uuid4().hex}.{extension}"
    file_storage.save(target)
    return str(Path("uploads") / folder / target.name).replace("\\", "/")


def _uploaded_path(relative_path):
    if not relative_path:
        return None
    path = (UPLOAD_ROOT.parent / relative_path).resolve()
    root = UPLOAD_ROOT.parent.resolve()
    if root not in path.parents and path != root:
        return None
    return path


def _passport_id(student_id):
    return f"IKC-SP-{datetime.now().year}-{int(student_id):05d}"


def _ensure_passports(cursor, student_id=None):
    query = "SELECT user_id FROM student_profiles WHERE passport_id IS NULL"
    params = ()
    if student_id is not None:
        query += " AND user_id = %s"
        params = (student_id,)
    cursor.execute(query, params)
    for row in cursor.fetchall():
        cursor.execute(
            "UPDATE student_profiles SET passport_id=%s WHERE user_id=%s AND passport_id IS NULL",
            (_passport_id(row["user_id"]), row["user_id"]),
        )


def _scope_for_role(role, user_id, alias="e.student_id"):
    if role == "student":
        return f"{alias} = %s", [user_id]
    if role == "coach":
        return (
            f"{alias} IN (SELECT ss.student_id FROM student_sports ss "
            "JOIN coach_sports cs ON cs.sport_id=ss.sport_id WHERE cs.coach_id=%s)",
            [user_id],
        )
    return "TRUE", []


def _projection(period):
    today = date.today()
    if period == "weekly":
        start = today - timedelta(days=6)
    elif period == "monthly":
        start = today.replace(day=1)
    else:
        period = "yearly"
        start = today.replace(month=1, day=1)
    return period, start.isoformat(), today.isoformat()


def _build_projection(cursor, period, role, user_id):
    period, start_date, end_date = _projection(period)
    scope, scope_params = _scope_for_role(role, user_id, "e.student_id")
    cursor.execute(
        f"""SELECT s.name AS sport,
                   COUNT(DISTINCT e.student_id) AS athletes,
                   COUNT(a.id) AS sessions,
                   COALESCE(SUM(CASE WHEN a.present THEN 1 ELSE 0 END), 0) AS present_sessions,
                   COALESCE(AVG(da.overall_performance), 0) AS avg_performance
            FROM programmes p
            JOIN sports s ON s.id=p.sport_id
            JOIN enrollments e ON e.programme_id=p.id AND e.status='Active'
            LEFT JOIN attendance a ON a.enrollment_id=e.id
                AND a.session_date BETWEEN %s AND %s
            LEFT JOIN daily_assessments da ON da.enrollment_id=e.id
                AND da.session_date BETWEEN %s AND %s
            WHERE p.active=TRUE AND {scope}
            GROUP BY s.id, s.name ORDER BY s.name""",
        (start_date, end_date, start_date, end_date, *scope_params),
    )
    cursor.execute(
        f"""SELECT s.name AS sport, ams.metric_key, ams.metric_label,
                   AVG(ams.score) AS avg_score, COUNT(*) AS samples
            FROM assessment_metric_scores ams
            JOIN daily_assessments da ON da.id=ams.assessment_id
            JOIN enrollments e ON e.id=da.enrollment_id
            JOIN programmes p ON p.id=e.programme_id
            JOIN sports s ON s.id=p.sport_id
            WHERE da.session_date BETWEEN %s AND %s AND {scope}
            GROUP BY s.name, ams.metric_key, ams.metric_label
            ORDER BY s.name, ams.metric_label""",
        (start_date, end_date, *scope_params),
    )
    metrics_by_sport = {}
    for metric in cursor.fetchall():
        metrics_by_sport.setdefault(metric["sport"], []).append(
            {
                "metric_key": metric["metric_key"],
                "metric_label": metric["metric_label"],
                "avg_score": float(metric["avg_score"] or 0),
                "samples": int(metric["samples"] or 0),
            }
        )
    sports = []
    for row in cursor.fetchall():
        sessions = int(row["sessions"] or 0)
        present = int(row["present_sessions"] or 0)
        sports.append(
            {
                "sport": row["sport"],
                "athletes": int(row["athletes"] or 0),
                "sessions": sessions,
                "present_sessions": present,
                "attendance_rate": round(present / sessions * 100, 1) if sessions else 0,
                "avg_performance": float(row["avg_performance"] or 0),
                "metrics": metrics_by_sport.get(row["sport"], []),
            }
        )
    return {"period": period, "from_date": start_date, "to_date": end_date, "sports": sports}


def _csv_response(filename, headers, rows):
    output = io.StringIO()
    output.write(",".join(headers) + "\n")
    for row in rows:
        output.write(",".join('"' + str(value or "").replace('"', '""') + '"' for value in row) + "\n")
    return Response(
        output.getvalue(),
        headers={"Content-Disposition": f"attachment; filename={filename}"},
        content_type="text/csv; charset=utf-8",
    )


def _participant_rows(upload):
    if not upload or not upload.filename or "." not in upload.filename:
        raise ValueError("Choose an XLSX or XLS participant file.")
    extension = upload.filename.rsplit(".", 1)[1].lower()
    raw_rows = []
    data = upload.read()
    if extension == "xlsx":
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise ValueError("XLSX support is not installed.") from exc
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        raw_rows = [list(row) for row in workbook.active.iter_rows(values_only=True)]
    elif extension == "xls":
        try:
            import xlrd
        except ImportError as exc:
            raise ValueError("XLS support is not installed.") from exc
        sheet = xlrd.open_workbook(file_contents=data).sheet_by_index(0)
        raw_rows = [sheet.row_values(index) for index in range(sheet.nrows)]
    else:
        raise ValueError("Participant import supports .xlsx and .xls files only.")
    if not raw_rows:
        raise ValueError("The participant sheet is empty.")
    headers = [str(value or "").strip().lower().replace("_", " ") for value in raw_rows[0]]
    aliases = {
        "name": {"name", "student name", "participant name"},
        "department": {"department", "dept"},
        "category": {"category", "gender category", "boys/girls"},
        "year": {"year", "academic year", "batch/year"},
    }
    positions = {}
    for key, options in aliases.items():
        for position, header in enumerate(headers):
            if header in options:
                positions[key] = position
                break
    missing = [key.title() for key in aliases if key not in positions]
    if missing:
        raise ValueError("Missing required participant columns: " + ", ".join(missing))
    result = []
    for raw in raw_rows[1:]:
        row = {
            key: str(raw[position] if position < len(raw) and raw[position] is not None else "").strip()
            for key, position in positions.items()
        }
        if not row["name"]:
            continue
        row["category"] = row["category"].title()
        if row["category"] not in {"Boys", "Girls"}:
            raise ValueError(f"Invalid category for {row['name']}: use Boys or Girls.")
        result.append(row)
    return result


def register_reference_extension_routes(app):
    @app.route("/api/fit-track/students/bulk-sports", methods=["POST"])
    @ADMIN
    def fit_track_bulk_sports():
        data = _body()
        student_ids = data.get("student_ids") or []
        sport_ids = data.get("sport_ids") or []
        if not student_ids or not sport_ids:
            return jsonify(success=False, message="student_ids and sport_ids are required."), 400
        db = get_db()
        try:
            with db.cursor() as cursor:
                for student_id in student_ids:
                    cursor.execute("SELECT 1 FROM users WHERE id=%s AND role='student'", (student_id,))
                    if not cursor.fetchone():
                        return jsonify(success=False, message=f"Student {student_id} was not found."), 404
                for sport_id in sport_ids:
                    cursor.execute("SELECT 1 FROM sports WHERE id=%s AND is_active=TRUE", (sport_id,))
                    if not cursor.fetchone():
                        return jsonify(success=False, message=f"Sport {sport_id} was not found."), 404
                for student_id in student_ids:
                    for sport_id in sport_ids:
                        cursor.execute(
                            "INSERT INTO student_sports(student_id,sport_id) VALUES(%s,%s) ON CONFLICT DO NOTHING",
                            (student_id, sport_id),
                        )
                _commit(db)
            return jsonify(success=True, assigned_students=len(student_ids), assigned_sports=len(sport_ids))
        finally:
            db.close()

    @app.route("/api/fit-track/enrollments/bulk", methods=["POST"])
    @STAFF
    def fit_track_bulk_enrollments():
        data = _body()
        programme_id = data.get("programme_id")
        student_ids = data.get("student_ids") or []
        if not programme_id or not student_ids:
            return jsonify(success=False, message="programme_id and student_ids are required."), 400
        db = get_db()
        try:
            with db.cursor() as cursor:
                if session.get("role") == "coach" and not _coach_can_access_programme(cursor, session.get("user_id"), programme_id):
                    return jsonify(success=False, message="You may only manage your assigned programme."), 403
                added = 0
                for student_id in student_ids:
                    cursor.execute("SELECT 1 FROM users WHERE id=%s AND role='student'", (student_id,))
                    if not cursor.fetchone():
                        return jsonify(success=False, message=f"Student {student_id} was not found."), 404
                    if session.get("role") == "coach" and not _coach_can_access_student(cursor, session.get("user_id"), student_id):
                        return jsonify(success=False, message=f"Student {student_id} is outside your assigned scope."), 403
                    cursor.execute(
                        """INSERT INTO enrollments(student_id,programme_id)
                           VALUES(%s,%s) ON CONFLICT(student_id,programme_id)
                           DO UPDATE SET status='Active'""",
                        (student_id, programme_id),
                    )
                    added += 1
                _commit(db)
            return jsonify(success=True, enrolled=added)
        finally:
            db.close()

    @app.route("/api/fit-track/reports/projection", methods=["GET"])
    @ALL_USERS
    def fit_track_projection():
        period = request.args.get("period", "yearly").lower()
        db = get_db()
        try:
            with db.cursor() as cursor:
                projection = _build_projection(cursor, period, session.get("role"), session.get("user_id"))
            return jsonify(success=True, data=projection)
        finally:
            db.close()

    @app.route("/api/fit-track/reports/projection.csv", methods=["GET"])
    @ALL_USERS
    def fit_track_projection_csv():
        period = request.args.get("period", "yearly").lower()
        db = get_db()
        try:
            with db.cursor() as cursor:
                projection = _build_projection(cursor, period, session.get("role"), session.get("user_id"))
            rows = []
            for sport in projection["sports"]:
                rows.append((projection["period"], projection["from_date"], projection["to_date"], sport["sport"], sport["athletes"], sport["sessions"], sport["present_sessions"], sport["attendance_rate"], sport["avg_performance"]))
            return _csv_response(
                f"fit-track-{projection['period']}-projection.csv",
                ["Period", "From", "To", "Sport", "Athletes", "Sessions", "Present Sessions", "Attendance Rate", "Average Performance"],
                rows,
            )
        finally:
            db.close()

    @app.route("/api/fit-track/students/<int:student_id>/performance.csv", methods=["GET"])
    @ALL_USERS
    def fit_track_student_performance_csv(student_id):
        db = get_db()
        try:
            with db.cursor() as cursor:
                if session.get("role") == "student" and session.get("user_id") != student_id:
                    return jsonify(success=False, message="You can only export your own performance."), 403
                if session.get("role") == "coach" and not _coach_can_access_student(cursor, session.get("user_id"), student_id):
                    return jsonify(success=False, message="That athlete is outside your assigned scope."), 403
                cursor.execute(
                    """SELECT da.session_date, s.name AS sport, p.name AS programme,
                              da.attendance_status, da.attendance_score,
                              da.overall_performance, da.remarks
                       FROM daily_assessments da
                       JOIN enrollments e ON e.id=da.enrollment_id
                       JOIN programmes p ON p.id=e.programme_id
                       JOIN sports s ON s.id=p.sport_id
                       WHERE e.student_id=%s ORDER BY da.session_date DESC""",
                    (student_id,),
                )
                rows = [
                    (r["session_date"], r["sport"], r["programme"], r["attendance_status"], r["attendance_score"], r["overall_performance"], r["remarks"])
                    for r in cursor.fetchall()
                ]
            return _csv_response("student-performance.csv", ["Date", "Sport", "Programme", "Attendance", "Attendance Score", "Overall Performance", "Remarks"], rows)
        finally:
            db.close()

    @app.route("/api/fit-track/tournaments/<int:tournament_id>/participants/import", methods=["POST"])
    @ADMIN
    def fit_track_import_participants(tournament_id):
        try:
            rows = _participant_rows(request.files.get("file"))
        except ValueError as error:
            return jsonify(success=False, message=str(error)), 400
        db = get_db()
        try:
            with db.cursor() as cursor:
                added = 0
                skipped = 0
                for participant in rows:
                    cursor.execute(
                        """SELECT u.id FROM users u JOIN student_profiles sp ON sp.user_id=u.id
                           WHERE u.role='student' AND LOWER(u.full_name)=LOWER(%s)
                             AND LOWER(COALESCE(sp.department,''))=LOWER(%s) LIMIT 1""",
                        (participant["name"], participant["department"]),
                    )
                    student = cursor.fetchone()
                    cursor.execute(
                        """INSERT INTO event_participants(tournament_id,student_id,name,department,category,year)
                           VALUES(%s,%s,%s,%s,%s,%s)
                           ON CONFLICT(tournament_id,name,department,category,year)
                           DO UPDATE SET student_id=EXCLUDED.student_id""",
                        (tournament_id, student["id"] if student else None, participant["name"], participant["department"], participant["category"], participant["year"]),
                    )
                    if cursor.rowcount:
                        added += 1
                    else:
                        skipped += 1
                _commit(db)
            return jsonify(success=True, added=added, skipped=skipped), 201
        finally:
            db.close()

    @app.route("/api/fit-track/participants/<int:participant_id>", methods=["DELETE"])
    @ADMIN
    def fit_track_delete_participant(participant_id):
        db = get_db()
        try:
            with db.cursor() as cursor:
                cursor.execute("DELETE FROM event_participants WHERE id=%s", (participant_id,))
                if cursor.rowcount == 0:
                    return jsonify(success=False, message="Participant not found."), 404
                _commit(db)
            return jsonify(success=True)
        finally:
            db.close()

    @app.route("/api/fit-track/matches/<int:match_id>", methods=["DELETE"])
    @ADMIN
    def fit_track_delete_match(match_id):
        db = get_db()
        try:
            with db.cursor() as cursor:
                cursor.execute("DELETE FROM matches WHERE id=%s", (match_id,))
                if cursor.rowcount == 0:
                    return jsonify(success=False, message="Match not found."), 404
                _commit(db)
            return jsonify(success=True)
        finally:
            db.close()

    @app.route("/api/fit-track/matches/<int:match_id>/players/<int:player_id>", methods=["DELETE"])
    @ADMIN
    def fit_track_delete_match_player(match_id, player_id):
        db = get_db()
        try:
            with db.cursor() as cursor:
                cursor.execute("DELETE FROM match_players WHERE id=%s AND match_id=%s", (player_id, match_id))
                if cursor.rowcount == 0:
                    return jsonify(success=False, message="Match player not found."), 404
                _commit(db)
            return jsonify(success=True)
        finally:
            db.close()

    @app.route("/api/fit-track/tournaments/<int:tournament_id>/results/<int:result_id>", methods=["DELETE"])
    @ADMIN
    def fit_track_delete_result(tournament_id, result_id):
        db = get_db()
        try:
            with db.cursor() as cursor:
                cursor.execute("DELETE FROM tournament_results WHERE id=%s AND tournament_id=%s", (result_id, tournament_id))
                if cursor.rowcount == 0:
                    return jsonify(success=False, message="Result not found."), 404
                _commit(db)
            return jsonify(success=True)
        finally:
            db.close()

    @app.route("/api/fit-track/draws/<int:draw_id>", methods=["DELETE"])
    @ADMIN
    def fit_track_delete_draw(draw_id):
        db = get_db()
        try:
            with db.cursor() as cursor:
                cursor.execute("DELETE FROM draw_documents WHERE id=%s", (draw_id,))
                if cursor.rowcount == 0:
                    return jsonify(success=False, message="Draw not found."), 404
                _commit(db)
            return jsonify(success=True)
        finally:
            db.close()

    @app.route("/api/fit-track/gallery/upload", methods=["POST"])
    @ADMIN
    def fit_track_upload_gallery():
        try:
            upload = request.files.get("file")
            allowed = IMAGE_EXTENSIONS | VIDEO_EXTENSIONS | DOCUMENT_EXTENSIONS
            path = _upload(upload, "events", allowed)
        except ValueError as error:
            return jsonify(success=False, message=str(error)), 400
        db = get_db()
        try:
            with db.cursor() as cursor:
                cursor.execute(
                    """INSERT INTO event_media(tournament_id,media_type,file_path,caption,uploaded_by)
                       VALUES(%s,%s,%s,%s,%s) RETURNING *""",
                    (request.form.get("tournament_id") or None, request.form.get("media_type") or "photo", path, request.form.get("caption"), session.get("user_id")),
                )
                row = _row(cursor.fetchone())
                _commit(db)
            return jsonify(success=True, data=row), 201
        finally:
            db.close()

    @app.route("/api/fit-track/announcements/upload", methods=["POST"])
    @ADMIN
    def fit_track_upload_announcement():
        title = (request.form.get("title") or "").strip()
        message = (request.form.get("message") or "").strip()
        if not title or not message:
            return jsonify(success=False, message="title and message are required."), 400
        try:
            image_path = _upload(request.files.get("image"), "announcements", IMAGE_EXTENSIONS) if request.files.get("image") else None
        except ValueError as error:
            return jsonify(success=False, message=str(error)), 400
        db = get_db()
        try:
            with db.cursor() as cursor:
                cursor.execute(
                    """INSERT INTO announcements(title,message,image_path,created_by)
                       VALUES(%s,%s,%s,%s) RETURNING *""",
                    (title, message, image_path, session.get("user_id")),
                )
                row = _row(cursor.fetchone())
                _commit(db)
            return jsonify(success=True, data=row), 201
        finally:
            db.close()

    @app.route("/api/fit-track/draws/upload", methods=["POST"])
    @ADMIN
    def fit_track_upload_draw():
        tournament_id = request.form.get("tournament_id")
        category = (request.form.get("category") or "Boys").title()
        format_type = (request.form.get("format_type") or "Knockout").title()
        if not tournament_id or category not in {"Boys", "Girls"} or format_type not in {"League", "Knockout"}:
            return jsonify(success=False, message="tournament_id, category, and format_type are required."), 400
        try:
            path = _upload(request.files.get("file"), "draws", DOCUMENT_EXTENSIONS)
        except ValueError as error:
            return jsonify(success=False, message=str(error)), 400
        db = get_db()
        try:
            with db.cursor() as cursor:
                cursor.execute(
                    """INSERT INTO draw_documents(tournament_id,category,format_type,file_path,uploaded_by)
                       VALUES(%s,%s,%s,%s,%s) RETURNING *""",
                    (tournament_id, category, format_type, path, session.get("user_id")),
                )
                row = _row(cursor.fetchone())
                _commit(db)
            return jsonify(success=True, data=row), 201
        finally:
            db.close()

    @app.route("/api/fit-track/students/<int:student_id>/photo", methods=["POST", "DELETE"])
    @ADMIN
    def fit_track_student_photo(student_id):
        db = get_db()
        try:
            with db.cursor() as cursor:
                if request.method == "DELETE":
                    cursor.execute("UPDATE student_profiles SET photo_path=NULL WHERE user_id=%s RETURNING user_id", (student_id,))
                    if not cursor.fetchone():
                        return jsonify(success=False, message="Student profile not found."), 404
                    _commit(db)
                    return jsonify(success=True)
                try:
                    path = _upload(request.files.get("photo"), "profiles/students", IMAGE_EXTENSIONS)
                except ValueError as error:
                    return jsonify(success=False, message=str(error)), 400
                cursor.execute("UPDATE student_profiles SET photo_path=%s WHERE user_id=%s RETURNING user_id,photo_path", (path, student_id))
                row = cursor.fetchone()
                if not row:
                    return jsonify(success=False, message="Student profile not found."), 404
                _commit(db)
            return jsonify(success=True, data=_row(row))
        finally:
            db.close()

    @app.route("/api/fit-track/coaches/<int:coach_id>/photo", methods=["POST", "DELETE"])
    @ADMIN
    def fit_track_coach_photo(coach_id):
        db = get_db()
        try:
            with db.cursor() as cursor:
                if request.method == "DELETE":
                    cursor.execute("UPDATE coach_profiles SET photo_path=NULL WHERE user_id=%s RETURNING user_id", (coach_id,))
                    if not cursor.fetchone():
                        return jsonify(success=False, message="Coach profile not found."), 404
                    _commit(db)
                    return jsonify(success=True)
                try:
                    path = _upload(request.files.get("photo"), "profiles/coaches", IMAGE_EXTENSIONS)
                except ValueError as error:
                    return jsonify(success=False, message=str(error)), 400
                cursor.execute(
                    """INSERT INTO coach_profiles(user_id,photo_path) VALUES(%s,%s)
                       ON CONFLICT(user_id) DO UPDATE SET photo_path=EXCLUDED.photo_path
                       RETURNING user_id,photo_path""",
                    (coach_id, path),
                )
                row = _row(cursor.fetchone())
                _commit(db)
            return jsonify(success=True, data=row)
        finally:
            db.close()

    @app.route("/api/fit-track/media/<int:media_id>/download", methods=["GET"])
    @ALL_USERS
    def fit_track_download_media(media_id):
        db = get_db()
        try:
            with db.cursor() as cursor:
                cursor.execute("SELECT file_path FROM event_media WHERE id=%s", (media_id,))
                row = cursor.fetchone()
            if not row:
                return jsonify(success=False, message="Media not found."), 404
        finally:
            db.close()
        path = _uploaded_path(row["file_path"])
        if not path or not path.is_file():
            return jsonify(success=False, message="Media file is unavailable."), 404
        return send_file(path, as_attachment=True)

    @app.route("/api/fit-track/facility-media/<int:media_id>/download", methods=["GET"])
    @ALL_USERS
    def fit_track_download_facility_media(media_id):
        db = get_db()
        try:
            with db.cursor() as cursor:
                cursor.execute("SELECT file_path FROM facility_media WHERE id=%s", (media_id,))
                row = cursor.fetchone()
            if not row:
                return jsonify(success=False, message="Facility media not found."), 404
        finally:
            db.close()
        path = _uploaded_path(row["file_path"])
        if not path or not path.is_file():
            return jsonify(success=False, message="Facility file is unavailable."), 404
        return send_file(path, as_attachment=True)

    @app.route("/api/fit-track/announcements/<int:announcement_id>/image", methods=["GET"])
    @ALL_USERS
    def fit_track_announcement_image(announcement_id):
        db = get_db()
        try:
            with db.cursor() as cursor:
                cursor.execute("SELECT image_path FROM announcements WHERE id=%s AND active=TRUE", (announcement_id,))
                row = cursor.fetchone()
            if not row:
                return jsonify(success=False, message="Announcement image not found."), 404
        finally:
            db.close()
        path = _uploaded_path(row["image_path"])
        if not path or not path.is_file():
            return jsonify(success=False, message="Announcement image is unavailable."), 404
        return send_file(path)

    @app.route("/api/fit-track/draws/<int:draw_id>/download", methods=["GET"])
    @ALL_USERS
    def fit_track_download_draw(draw_id):
        db = get_db()
        try:
            with db.cursor() as cursor:
                cursor.execute("SELECT file_path FROM draw_documents WHERE id=%s", (draw_id,))
                row = cursor.fetchone()
            if not row:
                return jsonify(success=False, message="Draw document not found."), 404
        finally:
            db.close()
        path = _uploaded_path(row["file_path"])
        if not path or not path.is_file():
            return jsonify(success=False, message="Draw document is unavailable."), 404
        return send_file(path, as_attachment=True)

    @app.route("/api/fit-track/facilities/<int:sport_id>/cover", methods=["POST", "DELETE"])
    @ADMIN
    def fit_track_facility_cover(sport_id):
        db = get_db()
        try:
            with db.cursor() as cursor:
                if request.method == "DELETE":
                    cursor.execute("UPDATE sports SET cover_path=NULL WHERE id=%s RETURNING id", (sport_id,))
                    if not cursor.fetchone():
                        return jsonify(success=False, message="Sport not found."), 404
                    _commit(db)
                    return jsonify(success=True)
                try:
                    path = _upload(request.files.get("file"), "facilities", IMAGE_EXTENSIONS)
                except ValueError as error:
                    return jsonify(success=False, message=str(error)), 400
                cursor.execute("UPDATE sports SET cover_path=%s WHERE id=%s RETURNING *", (path, sport_id))
                row = cursor.fetchone()
                if not row:
                    return jsonify(success=False, message="Sport not found."), 404
                _commit(db)
            return jsonify(success=True, data=_row(row))
        finally:
            db.close()

    @app.route("/api/fit-track/facilities/<int:sport_id>/media", methods=["POST"])
    @ADMIN
    def fit_track_upload_facility_media(sport_id):
        try:
            path = _upload(request.files.get("file"), "facilities", IMAGE_EXTENSIONS)
        except ValueError as error:
            return jsonify(success=False, message=str(error)), 400
        db = get_db()
        try:
            with db.cursor() as cursor:
                cursor.execute(
                    """INSERT INTO facility_media(sport_id,file_path,caption,uploaded_by)
                       VALUES(%s,%s,%s,%s) RETURNING *""",
                    (sport_id, path, request.form.get("caption"), session.get("user_id")),
                )
                row = _row(cursor.fetchone())
                _commit(db)
            return jsonify(success=True, data=row), 201
        finally:
            db.close()

    @app.route("/api/fit-track/certificate-templates/upload", methods=["POST"])
    @ADMIN
    def fit_track_upload_certificate_template():
        tournament_id = request.form.get("tournament_id")
        certificate_type = (request.form.get("certificate_type") or "Participation").title()
        if certificate_type not in {"Participation", "Achievement"} or not tournament_id:
            return jsonify(success=False, message="tournament_id and a valid certificate_type are required."), 400
        try:
            path = _upload(request.files.get("file"), "certificates/templates", CERTIFICATE_EXTENSIONS)
        except ValueError as error:
            return jsonify(success=False, message=str(error)), 400
        db = get_db()
        try:
            with db.cursor() as cursor:
                cursor.execute(
                    """INSERT INTO certificate_templates(tournament_id,certificate_type,file_path,uploaded_by)
                       VALUES(%s,%s,%s,%s)
                       ON CONFLICT(tournament_id,certificate_type)
                       DO UPDATE SET file_path=EXCLUDED.file_path, uploaded_by=EXCLUDED.uploaded_by
                       RETURNING *""",
                    (tournament_id, certificate_type, path, session.get("user_id")),
                )
                row = _row(cursor.fetchone())
                _commit(db)
            return jsonify(success=True, data=row), 201
        finally:
            db.close()

    @app.route("/api/fit-track/tournaments/<int:tournament_id>/certificates/generate-pdf", methods=["POST"])
    @ADMIN
    def fit_track_generate_certificates_pdf(tournament_id):
        data = _body()
        certificate_type = (data.get("certificate_type") or "Participation").title()
        if certificate_type not in {"Participation", "Achievement"}:
            return jsonify(success=False, message="Invalid certificate type."), 400
        db = get_db()
        try:
            with db.cursor() as cursor:
                cursor.execute("SELECT * FROM tournaments WHERE id=%s", (tournament_id,))
                event = cursor.fetchone()
                cursor.execute("SELECT * FROM certificate_templates WHERE tournament_id=%s AND certificate_type=%s", (tournament_id, certificate_type))
                template = cursor.fetchone()
                if not event or not template:
                    return jsonify(success=False, message="Tournament or certificate template not found."), 404
                template_path = _uploaded_path(template["file_path"])
                if not template_path or not template_path.is_file():
                    return jsonify(success=False, message="Certificate template file is unavailable."), 404
                cursor.execute("SELECT * FROM event_participants WHERE tournament_id=%s ORDER BY id", (tournament_id,))
                participants = cursor.fetchall()
                selected = [(participant, None) for participant in participants]
                if certificate_type == "Achievement":
                    cursor.execute("SELECT * FROM tournament_results WHERE tournament_id=%s", (tournament_id,))
                    results = cursor.fetchall()
                    selected = []
                    for participant in participants:
                        for result in results:
                            if participant["category"] != result["category"]:
                                continue
                            name = participant["name"].strip().lower()
                            department = participant["department"].strip().lower()
                            winner = result["winner"].strip().lower()
                            runner_up = result["runner_up"].strip().lower()
                            if name == winner or department == winner:
                                selected.append((participant, "Winner"))
                                break
                            if name == runner_up or department == runner_up:
                                selected.append((participant, "Runner-up"))
                                break
                if not selected:
                    return jsonify(success=False, message="No eligible participants found."), 400
                generated = []
                for participant, achievement in selected:
                    cursor.execute(
                        "SELECT id FROM generated_certificates WHERE tournament_id=%s AND participant_id=%s AND certificate_type=%s",
                        (tournament_id, participant["id"], certificate_type),
                    )
                    if cursor.fetchone():
                        continue
                    cursor.execute("SELECT COUNT(*) AS count FROM generated_certificates WHERE tournament_id=%s AND certificate_type=%s", (tournament_id, certificate_type))
                    serial = int(cursor.fetchone()["count"] or 0) + 1
                    prefix = "ACH" if certificate_type == "Achievement" else "PAR"
                    certificate_no = f"IKC-{datetime.now().year}-{int(tournament_id):03d}-{prefix}-{serial:04d}"
                    relative = _upload_generated_certificate(template_path, participant, event, certificate_type, certificate_no, achievement)
                    cursor.execute(
                        """INSERT INTO generated_certificates(tournament_id,participant_id,student_id,certificate_type,recipient_name,achievement,certificate_no,file_path)
                           VALUES(%s,%s,%s,%s,%s,%s,%s,%s) RETURNING *""",
                        (tournament_id, participant["id"], participant["student_id"], certificate_type, participant["name"], achievement, certificate_no, relative),
                    )
                    generated.append(_row(cursor.fetchone()))
                _commit(db)
            return jsonify(success=True, data=generated), 201
        except ValueError as error:
            db.rollback()
            return jsonify(success=False, message=str(error)), 400
        finally:
            db.close()

    @app.route("/api/fit-track/certificates/<int:certificate_id>/download", methods=["GET"])
    @ALL_USERS
    def fit_track_download_certificate(certificate_id):
        db = get_db()
        try:
            with db.cursor() as cursor:
                cursor.execute("SELECT * FROM generated_certificates WHERE id=%s", (certificate_id,))
                certificate = cursor.fetchone()
            if not certificate:
                return jsonify(success=False, message="Certificate not found."), 404
            if session.get("role") == "student" and certificate["student_id"] != session.get("user_id"):
                return jsonify(success=False, message="You may only download your own certificate."), 403
            if session.get("role") == "coach":
                return jsonify(success=False, message="Only the linked student or an admin may download this certificate."), 403
        finally:
            db.close()
        path = _uploaded_path(certificate["file_path"])
        if not path or not path.is_file():
            return jsonify(success=False, message="Certificate file is unavailable."), 404
        return send_file(path, as_attachment=True, download_name=f"{certificate['certificate_no']}.pdf")


def _upload_generated_certificate(template_path, participant, event, certificate_type, certificate_no, achievement):
    """Overlay certificate text using the same centered layout as the reference app."""
    try:
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.utils import ImageReader
        from reportlab.pdfgen import canvas
        from pypdf import PdfReader, PdfWriter
    except ImportError as exc:
        raise ValueError("Certificate generation dependencies are not installed.") from exc
    relative = Path("uploads") / "certificates" / "generated" / f"{uuid4().hex}.pdf"
    output_path = UPLOAD_ROOT.parent / relative
    output_path.parent.mkdir(parents=True, exist_ok=True)
    extension = template_path.suffix.lower()
    if extension == ".pdf":
        template = PdfReader(str(template_path))
        if not template.pages:
            raise ValueError("Certificate template PDF has no pages.")
        page = template.pages[0]
        width = float(page.mediabox.width)
        height = float(page.mediabox.height)
        overlay_stream = io.BytesIO()
        painter = canvas.Canvas(overlay_stream, pagesize=(width, height))
        painter.setFillColorRGB(0.02, 0.12, 0.35)
        painter.setFont("Helvetica-Bold", min(28, width / 25))
        painter.drawCentredString(width / 2, height * 0.49, participant["name"])
        painter.setFont("Helvetica", min(13, width / 55))
        painter.drawCentredString(width / 2, height * 0.41, f"{certificate_type} - {event['name']} - {participant['category']} - {participant['year']}")
        if achievement:
            painter.setFont("Helvetica-Bold", min(15, width / 48))
            painter.drawCentredString(width / 2, height * 0.35, achievement)
        painter.setFont("Helvetica", min(9, width / 75))
        painter.drawCentredString(width / 2, height * 0.10, f"Certificate No: {certificate_no}")
        painter.save()
        overlay_stream.seek(0)
        page.merge_page(PdfReader(overlay_stream).pages[0])
        writer = PdfWriter()
        writer.add_page(page)
        with output_path.open("wb") as output:
            writer.write(output)
    else:
        try:
            from PIL import Image
        except ImportError as exc:
            raise ValueError("Pillow is required for image certificate templates.") from exc
        image = Image.open(template_path)
        width_px, height_px = image.size
        width = landscape(A4)[0]
        height = width * (height_px / width_px)
        painter = canvas.Canvas(str(output_path), pagesize=(width, height))
        painter.drawImage(ImageReader(image), 0, 0, width=width, height=height, preserveAspectRatio=True, mask="auto")
        painter.setFillColorRGB(0.02, 0.12, 0.35)
        painter.setFont("Helvetica-Bold", 26)
        painter.drawCentredString(width / 2, height * 0.49, participant["name"])
        painter.setFont("Helvetica", 12)
        painter.drawCentredString(width / 2, height * 0.41, f"{certificate_type} - {event['name']} - {participant['category']} - {participant['year']}")
        if achievement:
            painter.setFont("Helvetica-Bold", 14)
            painter.drawCentredString(width / 2, height * 0.35, achievement)
        painter.setFont("Helvetica", 8)
        painter.drawCentredString(width / 2, height * 0.09, f"Certificate No: {certificate_no}")
        painter.save()
    return str(relative).replace("\\", "/")
